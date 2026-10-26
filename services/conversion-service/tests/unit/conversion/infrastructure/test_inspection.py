"""The scanned-PDF check that sends a file to OCR, and the page counts reported as usage."""

from pathlib import Path

import pikepdf
import pptx
from pikepdf import Dictionary, Name

from conversion.conversion.domain.formats import Format
from conversion.conversion.infrastructure.inspection import TEXT_CHECK_PAGES, PdfiumFileInspector


def _pdf(path: Path, pages_with_text: set[int], *, pages: int = 1) -> Path:
    pdf = pikepdf.new()
    font = Dictionary(Type=Name.Font, Subtype=Name.Type1, BaseFont=Name("/Helvetica"))
    for index in range(pages):
        page = pdf.add_blank_page(page_size=(612, 792))
        page.Resources = Dictionary(Font=Dictionary(F1=font))
        page.Contents = pdf.make_stream(b"BT /F1 12 Tf 72 700 Td (Hello) Tj ET" if index in pages_with_text else b"")
    pdf.save(path)
    return path


def test_treats_a_pdf_with_text_on_an_early_page_as_digital(tmp_path: Path) -> None:
    assert not PdfiumFileInspector().is_scanned_pdf(_pdf(tmp_path / "a.pdf", {2}, pages=3))


def test_treats_a_pdf_with_no_text_in_its_first_pages_as_scanned(tmp_path: Path, make_pdf) -> None:
    late_text = _pdf(tmp_path / "late.pdf", {TEXT_CHECK_PAGES}, pages=TEXT_CHECK_PAGES + 1)

    assert PdfiumFileInspector().is_scanned_pdf(make_pdf("scan.pdf", sizes=[(200, 100)]))
    assert PdfiumFileInspector().is_scanned_pdf(late_text)


def test_counts_pages_of_pdfs_slides_and_images_but_not_sheets(tmp_path: Path) -> None:
    deck = tmp_path / "deck.pptx"
    presentation = pptx.Presentation()
    for _ in range(2):
        presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(str(deck))
    inspector = PdfiumFileInspector()

    assert inspector.count_pages(_pdf(tmp_path / "p.pdf", set(), pages=3), Format.PDF) == 3
    assert inspector.count_pages(deck, Format.PPTX) == 2
    assert inspector.count_pages(tmp_path / "x.png", Format.PNG) == 1
    assert inspector.count_pages(tmp_path / "x.csv", Format.CSV) is None
