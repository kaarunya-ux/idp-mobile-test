"""Word -> model: roles from styles, alignment and bold kept, merged cells read once, header rows found."""

from collections.abc import Callable
from pathlib import Path

import docx
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from conversion.conversion.domain.document_model import Alignment, DocumentModel, Table, TextBlock, TextKind
from conversion.conversion.domain.exceptions import UnreadableDocumentError
from conversion.conversion.domain.formats import Representation
from conversion.conversion.domain.ports import StepRequest
from conversion.conversion.domain.route import Converter
from conversion.conversion.infrastructure.converters.docx_to_model import PythonDocxReader
from conversion.conversion.infrastructure.model_file import read_model

R = Representation
MakeRequest = Callable[..., StepRequest]
Word = docx.document.Document


def _read(document: Word, tmp_path: Path, make_request: MakeRequest) -> DocumentModel:
    source = tmp_path / "doc.docx"
    document.save(str(source))
    (output,) = (
        PythonDocxReader().convert(make_request(Converter.DOCX_TO_MODEL, R.DOCX, R.DOCUMENT_MODEL, [source])).outputs
    )
    return read_model(output)


def _texts(model: DocumentModel) -> list[TextBlock]:
    return [block for block in model.pages[0].blocks if isinstance(block, TextBlock)]


def test_reads_heading_list_caption_and_paragraph_roles_from_styles(tmp_path: Path, make_request: MakeRequest) -> None:
    document = docx.Document()
    document.add_heading("Report", 0)  # Word's "Title" style
    document.add_heading("Section", 2)
    document.add_heading("Deep", 9)
    document.add_paragraph("First point", style="List Bullet")
    document.add_paragraph("Sub point", style="List Bullet 2")
    document.add_paragraph("Figure 1", style="Caption")
    document.add_paragraph("Body text")
    document.add_paragraph("   ")  # empty paragraphs are skipped

    roles = [(t.kind, t.level, t.text) for t in _texts(_read(document, tmp_path, make_request))]

    assert roles == [
        (TextKind.HEADING, 1, "Report"),
        (TextKind.HEADING, 2, "Section"),
        (TextKind.HEADING, 6, "Deep"),  # Word goes to 9; the model caps at 6
        (TextKind.LIST_ITEM, 1, "First point"),
        (TextKind.LIST_ITEM, 2, "Sub point"),
        (TextKind.CAPTION, 1, "Figure 1"),
        (TextKind.PARAGRAPH, 1, "Body text"),
    ]


def test_reads_list_depth_from_numbering_when_the_style_does_not_say(tmp_path: Path, make_request: MakeRequest) -> None:
    document = docx.Document()
    paragraph = document.add_paragraph("Nested item")
    numbering = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "2")
    numbering.append(level)
    paragraph._p.get_or_add_pPr().append(numbering)

    (block,) = _texts(_read(document, tmp_path, make_request))

    assert (block.kind, block.level) == (TextKind.LIST_ITEM, 3)


def test_keeps_alignment_and_bold_when_the_paragraph_has_them(tmp_path: Path, make_request: MakeRequest) -> None:
    document = docx.Document()
    centred = document.add_paragraph()
    centred.add_run("Centred and bold").bold = True
    centred.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("Right").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    document.add_paragraph("Justified").alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    blocks = _texts(_read(document, tmp_path, make_request))

    assert [(b.alignment, b.bold) for b in blocks] == [
        (Alignment.CENTER, True),
        (Alignment.RIGHT, False),
        (Alignment.LEFT, False),  # justified text reads as left-aligned
    ]


def test_reads_a_merged_cell_once_when_it_spans_rows_or_columns(tmp_path: Path, make_request: MakeRequest) -> None:
    document = docx.Document()
    grid = document.add_table(rows=3, cols=3)
    grid.cell(0, 0).merge(grid.cell(0, 1)).text = "Wide"
    grid.cell(1, 2).merge(grid.cell(2, 2)).text = "Tall"
    grid.cell(0, 2).text = "C"

    (table,) = _read(document, tmp_path, make_request).tables

    assert table.rows == (("Wide", "", "C"), ("", "", "Tall"), ("", "", ""))


def test_counts_rows_word_repeats_as_headers_when_they_are_marked(tmp_path: Path, make_request: MakeRequest) -> None:
    document = docx.Document()
    grid = document.add_table(rows=3, cols=1)
    for index, text in enumerate(["Header", "Sub-header", "Data"]):
        grid.cell(index, 0).text = text
    for row in grid.rows[:2]:
        row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

    (table,) = _read(document, tmp_path, make_request).tables

    assert table.header_rows == 2


@pytest.mark.parametrize(("bold_first_row", "expected"), [(True, 1), (False, 0)])
def test_treats_a_bold_first_row_as_a_header_when_nothing_is_marked(
    tmp_path: Path, make_request: MakeRequest, bold_first_row: bool, expected: int
) -> None:
    document = docx.Document()
    grid = document.add_table(rows=2, cols=2)
    for column, text in enumerate(["item", "qty"]):
        grid.cell(0, column).paragraphs[0].add_run(text).bold = bold_first_row
        grid.cell(1, column).text = "x"

    (table,) = _read(document, tmp_path, make_request).tables

    assert table.header_rows == expected


def test_skips_a_table_when_every_cell_is_empty(tmp_path: Path, make_request: MakeRequest) -> None:
    document = docx.Document()
    document.add_table(rows=2, cols=2)
    document.add_paragraph("After")

    model = _read(document, tmp_path, make_request)

    assert model.tables == []
    assert [b.text for b in _texts(model)] == ["After"]


def test_reports_cnv_422_file_001_when_the_file_is_not_a_word_document(
    tmp_path: Path, make_request: MakeRequest
) -> None:
    fake = tmp_path / "fake.docx"
    fake.write_bytes(b"not a zip at all")

    with pytest.raises(UnreadableDocumentError):
        PythonDocxReader().convert(make_request(Converter.DOCX_TO_MODEL, R.DOCX, R.DOCUMENT_MODEL, [fake]))


def test_keeps_tables_and_text_in_document_order(tmp_path: Path, make_request: MakeRequest) -> None:
    document = docx.Document()
    document.add_paragraph("Before")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Cell"
    document.add_paragraph("After")

    blocks = _read(document, tmp_path, make_request).pages[0].blocks

    assert [type(block).__name__ for block in blocks] == ["TextBlock", "Table", "TextBlock"]
    assert isinstance(blocks[1], Table)
