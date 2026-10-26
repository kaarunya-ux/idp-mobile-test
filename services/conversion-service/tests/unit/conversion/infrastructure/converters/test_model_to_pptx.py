"""Model -> PowerPoint: titled slides, continuation slides, tables of at most 15 rows, real bullets."""

from collections.abc import Callable
from pathlib import Path

import pptx
import pytest
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.presentation import Presentation
from pptx.shapes.autoshape import Shape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.slide import Slide
from pptx.util import Pt

from conversion.conversion.domain.document_model import (
    Alignment,
    DocumentModel,
    Page,
    Table,
    TextBlock,
    TextKind,
)
from conversion.conversion.domain.exceptions import LimitExceededError
from conversion.conversion.domain.formats import Representation
from conversion.conversion.domain.ports import StepRequest
from conversion.conversion.domain.route import Converter
from conversion.conversion.infrastructure.converters.model_to_pptx import (
    ROWS_PER_SLIDE,
    TEXT_BLOCKS_PER_SLIDE,
    PythonPptxWriter,
)
from conversion.conversion.infrastructure.converters.pptx_to_model import PythonPptxReader
from conversion.conversion.infrastructure.model_file import read_model, write_model

R = Representation
MakeRequest = Callable[..., StepRequest]
WIDESCREEN = (12_192_000, 6_858_000)


def _write(tmp_path: Path, model: DocumentModel, make_request: MakeRequest, max_rows: int = 1000) -> Path:
    model_file = tmp_path / "report.model.json"
    write_model(model, model_file)
    request = make_request(Converter.MODEL_TO_PPTX, R.DOCUMENT_MODEL, R.PPTX, [model_file])
    (output,) = PythonPptxWriter(max_table_rows=max_rows).convert(request).outputs
    return output


def _deck(tmp_path: Path, model: DocumentModel, make_request: MakeRequest) -> Presentation:
    return pptx.Presentation(str(_write(tmp_path, model, make_request)))


def _titles(deck: Presentation) -> list[str]:
    return [slide.shapes.title.text_frame.text for slide in deck.slides if slide.shapes.title is not None]


def _tables(slide: Slide) -> list[GraphicFrame]:
    return [shape for shape in slide.shapes if isinstance(shape, GraphicFrame) and shape.has_table]


def _text_box(slide: Slide) -> Shape:
    (box,) = [shape for shape in slide.shapes if isinstance(shape, Shape) and not shape.is_placeholder]
    return box


def _rows(n: int) -> tuple[tuple[str, ...], ...]:
    return tuple((f"item {i}", str(i)) for i in range(n))


def test_writes_a_widescreen_deck_titled_after_each_page(tmp_path: Path, make_request: MakeRequest) -> None:
    model = DocumentModel(pages=(Page(name="Sales", blocks=()), Page(name="Staff", blocks=())))

    output = _write(tmp_path, model, make_request)

    deck = pptx.Presentation(str(output))
    assert output.name == "report.pptx"
    assert (deck.slide_width, deck.slide_height) == WIDESCREEN
    assert _titles(deck) == ["Sales", "Staff"]


def test_writes_each_text_role_with_its_look(tmp_path: Path, make_request: MakeRequest) -> None:
    model = DocumentModel(
        pages=(
            Page(
                name="Notes",
                blocks=(
                    TextBlock(text="Summary", kind=TextKind.HEADING, level=2),
                    TextBlock(text="Point", kind=TextKind.LIST_ITEM, level=2),
                    TextBlock(text="Figure 1", kind=TextKind.CAPTION),
                    TextBlock(text="Right", alignment=Alignment.RIGHT),
                ),
            ),
        )
    )

    heading, point, caption, right = _text_box(_deck(tmp_path, model, make_request).slides[0]).text_frame.paragraphs

    assert (heading.runs[0].font.bold, heading.runs[0].font.size) == (True, Pt(22))
    assert point.level == 1
    assert point._p.pPr is not None
    assert point._p.pPr.find(qn("a:buChar")) is not None  # a real bullet, not a typed character
    assert caption.runs[0].font.italic
    assert right.alignment == PP_ALIGN.RIGHT


def test_continues_a_long_table_on_slides_that_repeat_its_header(tmp_path: Path, make_request: MakeRequest) -> None:
    table = Table(rows=(("item", "qty"), *_rows(30)), header_rows=1)

    deck = _deck(tmp_path, DocumentModel(pages=(Page(name="Sales", blocks=(table,)),)), make_request)

    assert _titles(deck) == ["Sales", "Sales (cont.)", "Sales (cont.)"]
    parts = [_tables(slide)[0].table for slide in deck.slides]
    assert [len(part.rows) for part in parts] == [ROWS_PER_SLIDE, ROWS_PER_SLIDE, 3]  # 14 + 14 + 2 data rows
    assert all(part.cell(0, 0).text == "item" and part.first_row for part in parts)
    assert parts[1].cell(1, 0).text == "item 14"


def test_treats_a_header_too_tall_for_a_slide_as_data(tmp_path: Path, make_request: MakeRequest) -> None:
    table = Table(rows=_rows(ROWS_PER_SLIDE + 1), header_rows=ROWS_PER_SLIDE)

    deck = _deck(tmp_path, DocumentModel(pages=(Page(name="Odd", blocks=(table,)),)), make_request)

    parts = [_tables(slide)[0].table for slide in deck.slides]
    assert [len(part.rows) for part in parts] == [ROWS_PER_SLIDE, 1]
    assert not parts[0].first_row


def test_puts_text_and_tables_on_their_own_slides_in_order(tmp_path: Path, make_request: MakeRequest) -> None:
    texts = tuple(TextBlock(text=f"line {i}") for i in range(TEXT_BLOCKS_PER_SLIDE + 1))
    table = Table(rows=(("a",), ("b",)), header_rows=1)

    deck = _deck(tmp_path, DocumentModel(pages=(Page(name="Mix", blocks=(*texts, table, texts[0])),)), make_request)

    slides = list(deck.slides)
    assert len(slides) == 4  # 10 lines, 1 line, the table, 1 line
    assert len(_text_box(slides[0]).text_frame.paragraphs) == TEXT_BLOCKS_PER_SLIDE
    assert _tables(slides[2])


@pytest.mark.parametrize(
    ("blocks", "expected_title", "expected_text"),
    [
        ((TextBlock(text="Report", kind=TextKind.HEADING), TextBlock(text="Body")), "Report", ["Body"]),
        ((TextBlock(text="Body"),), "report", ["Body"]),  # no name, no heading: the file name
    ],
)
def test_titles_an_unnamed_page_with_its_first_heading_or_the_file_name(
    tmp_path: Path,
    make_request: MakeRequest,
    blocks: tuple[TextBlock, ...],
    expected_title: str,
    expected_text: list[str],
) -> None:
    deck = _deck(tmp_path, DocumentModel(pages=(Page(blocks=blocks),)), make_request)

    assert _titles(deck) == [expected_title]
    assert [p.text for p in _text_box(deck.slides[0]).text_frame.paragraphs] == expected_text


def test_writes_one_titled_slide_when_the_model_is_empty(tmp_path: Path, make_request: MakeRequest) -> None:
    assert _titles(_deck(tmp_path, DocumentModel(pages=()), make_request)) == ["report"]


@pytest.mark.parametrize(("columns", "size"), [(2, Pt(14)), (8, Pt(12)), (12, Pt(10)), (20, Pt(8))])
def test_shrinks_cell_text_as_tables_get_wider(
    tmp_path: Path, make_request: MakeRequest, columns: int, size: Pt
) -> None:
    table = Table(rows=(tuple("x" * columns),))

    deck = _deck(tmp_path, DocumentModel(pages=(Page(name="Wide", blocks=(table,)),)), make_request)

    assert _tables(deck.slides[0])[0].table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.size == size


def test_refuses_a_table_too_long_for_a_deck(tmp_path: Path, make_request: MakeRequest) -> None:
    long_table = Table(rows=_rows(201))

    with pytest.raises(LimitExceededError, match="XLSX or CSV"):
        _write(tmp_path, DocumentModel(pages=(Page(blocks=(long_table,)),)), make_request, max_rows=200)


def test_reads_back_the_same_content_when_the_deck_is_read_again(tmp_path: Path, make_request: MakeRequest) -> None:
    page = Page(
        name="Sales",
        blocks=(
            TextBlock(text="Up 10%", kind=TextKind.LIST_ITEM, level=2),
            Table(rows=(("item", "qty"), ("Paper", "10")), header_rows=1),
        ),
    )
    deck = _write(tmp_path, DocumentModel(pages=(page,)), make_request)

    request = make_request(Converter.PPTX_TO_MODEL, R.PPTX, R.DOCUMENT_MODEL, [deck])
    (model_file,) = PythonPptxReader().convert(request).outputs

    pages = read_model(model_file).pages
    assert [p.name for p in pages] == ["Sales", "Sales (cont.)"]  # text and the table get a slide each
    assert tuple(block for p in pages for block in p.blocks) == page.blocks
