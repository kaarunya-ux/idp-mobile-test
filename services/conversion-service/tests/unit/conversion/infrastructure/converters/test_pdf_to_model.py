"""Digital PDF -> model: headings by size, paragraphs from lines, bullets, ruled tables, PDF errors, full routes."""

import shutil
from collections.abc import Callable
from datetime import timedelta
from pathlib import Path
from typing import Self

import docx
import openpyxl
import pikepdf
import pytest
from pdfplumber.utils.exceptions import PdfminerException
from pikepdf import Dictionary, Name

from conversion.bootstrap import build_convert_file_service
from conversion.conversion.application.commands import ConvertFileCommand
from conversion.conversion.domain.document_model import Alignment, DocumentModel, Table, TextBlock, TextKind
from conversion.conversion.domain.exceptions import CorruptPdfError, LimitExceededError, PasswordRequiredError
from conversion.conversion.domain.formats import Format, Representation
from conversion.conversion.domain.options import ConversionOptions, PageRange
from conversion.conversion.domain.ports import StepRequest
from conversion.conversion.domain.route import Converter
from conversion.conversion.infrastructure.converters import pdf_to_model
from conversion.conversion.infrastructure.converters.office_to_pdf import LibreOfficePdfConverter
from conversion.conversion.infrastructure.converters.pdf_to_model import PdfplumberReader
from conversion.conversion.infrastructure.model_file import read_model
from conversion.core.config import Settings

R = Representation
MakeRequest = Callable[..., StepRequest]
MakePdf = Callable[..., Path]
PAGE_WIDTH, PAGE_HEIGHT = 612, 792  # US Letter, in points


class PdfBuilder:
    """Draws text in Helvetica and ruled grids on US Letter pages, measured from the top-left like pdfplumber."""

    def __init__(self) -> None:
        self._pages: list[bytes] = [b""]

    def text(self, x: float, top: float, text: str, *, size: int = 11, bold: bool = False) -> Self:
        escaped = text.encode("cp1252").replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")
        font = b"F2" if bold else b"F1"
        self._pages[-1] += b"BT /%s %d Tf %.1f %.1f Td (%s) Tj ET\n" % (
            font,
            size,
            x,
            PAGE_HEIGHT - top - size,
            escaped,
        )
        return self

    def grid(self, x: float, top: float, rows: list[list[str]], *, width: int = 120, height: int = 20) -> Self:
        for row_index, row in enumerate(rows):
            for column, value in enumerate(row):
                left, cell_top = x + column * width, top + row_index * height
                self._pages[-1] += b"%.1f %.1f %d %d re S\n" % (left, PAGE_HEIGHT - cell_top - height, width, height)
                self.text(left + 4, cell_top + 5, value)
        return self

    def page(self) -> Self:
        self._pages.append(b"")
        return self

    def save(self, path: Path) -> Path:
        pdf = pikepdf.new()
        fonts = Dictionary(F1=_font("/Helvetica"), F2=_font("/Helvetica-Bold"))
        for content in self._pages:
            page = pdf.add_blank_page(page_size=(PAGE_WIDTH, PAGE_HEIGHT))
            page.Resources = Dictionary(Font=fonts)
            page.Contents = pdf.make_stream(content)
        pdf.save(path)
        return path


def _font(base: str) -> Dictionary:
    return Dictionary(Type=Name.Font, Subtype=Name.Type1, BaseFont=Name(base), Encoding=Name.WinAnsiEncoding)


def _read(
    pdf: PdfBuilder, tmp_path: Path, make_request: MakeRequest, *, max_pages: int = 100, **options: str
) -> DocumentModel:
    source = pdf.save(tmp_path / "scan.pdf")
    request = make_request(Converter.PDF_TO_MODEL, R.PDF, R.DOCUMENT_MODEL, [source], ConversionOptions(**options))
    (output,) = PdfplumberReader(max_pages=max_pages).convert(request).outputs
    return read_model(output)


def _texts(model: DocumentModel, page: int = 0) -> list[TextBlock]:
    return [block for block in model.pages[page].blocks if isinstance(block, TextBlock)]


def test_finds_headings_by_size_and_merges_lines_into_paragraphs(tmp_path: Path, make_request: MakeRequest) -> None:
    pdf = (
        PdfBuilder()
        .text(72, 72, "Annual Report", size=20, bold=True)
        .text(72, 110, "Summary", size=14)
        .text(72, 140, "This paragraph starts on one line")
        .text(72, 154, "and carries on to the next.")
        .text(72, 190, "A second paragraph, further down.")
    )

    blocks = _texts(_read(pdf, tmp_path, make_request))

    assert blocks == [
        TextBlock(text="Annual Report", kind=TextKind.HEADING, level=1, bold=True),
        TextBlock(text="Summary", kind=TextKind.HEADING, level=2),
        TextBlock(text="This paragraph starts on one line and carries on to the next."),
        TextBlock(text="A second paragraph, further down."),
    ]


def test_reads_bullets_and_numbers_as_list_items_with_their_depth(tmp_path: Path, make_request: MakeRequest) -> None:
    pdf = (
        PdfBuilder()
        .text(90, 72, "\u2022 First point")  # first-level bullets are indented in Word too
        .text(108, 86, "- Nested point")
        .text(90, 100, "2. Numbered point")
        .text(72, 140, "Body text to set the body size, long enough to count the most.")
    )

    items = [(b.kind, b.level, b.text) for b in _texts(_read(pdf, tmp_path, make_request))][:3]

    assert items == [
        (TextKind.LIST_ITEM, 1, "First point"),  # the bullet glyph goes: Word draws its own
        (TextKind.LIST_ITEM, 2, "Nested point"),
        (TextKind.LIST_ITEM, 1, "2. Numbered point"),  # numbers stay: they carry meaning
    ]


def test_keeps_centred_and_right_aligned_lines(tmp_path: Path, make_request: MakeRequest) -> None:
    pdf = (
        PdfBuilder()
        .text(72, 72, "A long left-aligned line, wider than a centred title would be on this page.")
        .text(280, 120, "Centred")
        .text(470, 160, "Right edge text")  # ends furthest right, so it sets the text's right margin
        .text(72, 200, "Short left")
    )

    alignments = [b.alignment for b in _texts(_read(pdf, tmp_path, make_request))]

    assert alignments == [Alignment.LEFT, Alignment.CENTER, Alignment.RIGHT, Alignment.LEFT]


def test_reads_a_ruled_table_in_place_without_repeating_its_text(tmp_path: Path, make_request: MakeRequest) -> None:
    pdf = (
        PdfBuilder()
        .text(72, 72, "Before the table.")
        .grid(72, 100, [["item", "qty"], ["Paper", "10"], ["Ink", "2"]])
        .text(72, 200, "After the table.")
    )

    blocks = _read(pdf, tmp_path, make_request).pages[0].blocks

    assert blocks == (
        TextBlock(text="Before the table."),
        Table(rows=(("item", "qty"), ("Paper", "10"), ("Ink", "2")), header_rows=1),
        TextBlock(text="After the table."),
    )


@pytest.mark.parametrize(
    ("rows", "expected_text"),
    [([["Important"], ["notice"]], ["Important notice"]), ([["", ""], ["", ""]], [])],
)
def test_reads_a_one_column_frame_as_text_and_skips_an_empty_grid(
    tmp_path: Path, make_request: MakeRequest, rows: list[list[str]], expected_text: list[str]
) -> None:
    model = _read(PdfBuilder().grid(72, 72, rows, width=200, height=14), tmp_path, make_request)

    assert model.tables == []
    assert [b.text for b in _texts(model)] == expected_text


def test_reads_only_the_selected_pages_when_a_page_range_is_given(tmp_path: Path, make_request: MakeRequest) -> None:
    pdf = PdfBuilder().text(72, 72, "One").page().text(72, 72, "Two").page()  # the third page is empty

    model = _read(pdf, tmp_path, make_request, page_range=PageRange.parse("2-3"))

    assert [[b.text for b in _texts(model, page)] for page in range(2)] == [["Two"], []]


def test_refuses_a_pdf_with_more_pages_than_the_limit(tmp_path: Path, make_request: MakeRequest) -> None:
    with pytest.raises(LimitExceededError, match="more than 1 pages"):
        _read(PdfBuilder().page(), tmp_path, make_request, max_pages=1)


def test_reports_the_pdf_errors_when_the_file_is_locked_or_damaged(
    tmp_path: Path, make_request: MakeRequest, make_pdf: MakePdf
) -> None:
    locked = make_pdf("locked.pdf", sizes=[(100, 100)], password="secret")
    damaged = tmp_path / "damaged.pdf"
    damaged.write_bytes(b"%PDF-1.7 not really")
    reader = PdfplumberReader(max_pages=10)

    for source, error in [(locked, PasswordRequiredError), (damaged, CorruptPdfError)]:
        with pytest.raises(error):
            reader.convert(make_request(Converter.PDF_TO_MODEL, R.PDF, R.DOCUMENT_MODEL, [source]))


def test_reports_a_damaged_pdf_when_pdfminer_cannot_parse_what_pdfium_repaired(
    tmp_path: Path, make_request: MakeRequest, monkeypatch: pytest.MonkeyPatch
) -> None:
    def refuse(_: Path) -> None:
        raise PdfminerException("unparsable")

    monkeypatch.setattr(pdf_to_model.pdfplumber, "open", refuse)

    with pytest.raises(CorruptPdfError):
        _read(PdfBuilder().text(72, 72, "x"), tmp_path, make_request)


def test_turns_a_digital_pdf_into_word_and_excel_through_the_model(tmp_path: Path) -> None:
    source = (
        PdfBuilder()
        .text(72, 72, "Invoice", size=18)
        .grid(72, 110, [["item", "qty"], ["Paper", "10"]])
        .text(72, 180, "Thank you.")
        .save(tmp_path / "invoice.pdf")
    )
    service = build_convert_file_service(Settings(_env_file=None))

    word = service.convert(
        ConvertFileCommand(source=Format.PDF, target=Format.DOCX, inputs=(source,), output_dir=tmp_path / "w")
    )
    excel = service.convert(
        ConvertFileCommand(source=Format.PDF, target=Format.XLSX, inputs=(source,), output_dir=tmp_path / "x")
    )

    document = docx.Document(str(word.outputs[0]))
    assert [(p.style.name, p.text) for p in document.paragraphs if p.style is not None] == [
        ("Heading 1", "Invoice"),
        ("Normal", "Thank you."),
    ]
    assert [[c.text for c in row.cells] for row in document.tables[0].rows] == [["item", "qty"], ["Paper", "10"]]
    sheet = openpyxl.load_workbook(excel.outputs[0]).active
    assert sheet is not None
    assert [[c.value for c in row] for row in sheet.iter_rows()] == [["item", "qty"], ["Paper", 10]]
    assert word.route.converters == (Converter.PDF_TO_MODEL, Converter.MODEL_TO_DOCX)


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice is not installed")
def test_reads_word_bullets_and_headings_back_from_a_pdf_libreoffice_made(
    tmp_path: Path, make_request: MakeRequest
) -> None:
    document = docx.Document()
    document.add_heading("Report", 1)
    document.add_paragraph("Exports", style="List Bullet")  # drawn with Symbol's private-area bullet, U+F0B7
    document.add_paragraph("Online", style="List Bullet 2")
    document.add_paragraph("Body text, long enough to be the most common size on the page.")
    source = tmp_path / "report.docx"
    document.save(str(source))
    office = LibreOfficePdfConverter(binary="soffice", timeout=timedelta(seconds=60))
    (pdf,) = office.convert(make_request(Converter.OFFICE_TO_PDF, R.DOCX, R.PDF, [source])).outputs

    (output,) = (
        PdfplumberReader(max_pages=10)
        .convert(make_request(Converter.PDF_TO_MODEL, R.PDF, R.DOCUMENT_MODEL, [pdf]))
        .outputs
    )

    assert [(b.kind, b.level, b.text) for b in _texts(read_model(output))][:3] == [
        (TextKind.HEADING, 1, "Report"),
        (TextKind.LIST_ITEM, 1, "Exports"),
        (TextKind.LIST_ITEM, 2, "Online"),
    ]
