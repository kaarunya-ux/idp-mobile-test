"""XLSX or CSV -> model: one named page per sheet with data, trimmed, header detected."""

from collections.abc import Callable
from pathlib import Path

import openpyxl

from conversion.conversion.domain.document_model import DocumentModel
from conversion.conversion.domain.formats import Representation
from conversion.conversion.domain.ports import StepRequest
from conversion.conversion.domain.route import Converter
from conversion.conversion.infrastructure.converters.sheet_to_model import SheetReader
from conversion.conversion.infrastructure.model_file import read_model

R = Representation
MakeRequest = Callable[..., StepRequest]


def _read(source: Path, kind: R, make_request: MakeRequest) -> DocumentModel:
    request = make_request(Converter.SHEET_TO_MODEL, kind, R.DOCUMENT_MODEL, [source])
    (output,) = SheetReader().convert(request).outputs
    return read_model(output)


def test_makes_one_named_page_per_sheet_with_data_when_reading_a_workbook(
    tmp_path: Path, make_request: MakeRequest
) -> None:
    book = openpyxl.Workbook()
    sales = book.active
    assert sales is not None
    sales.title = "Sales"
    sales.append(["item", "qty"])
    sales.append(["Paper", 10])
    sales["D9"].number_format = "0.00"  # a styled empty cell: trailing rows and columns are trimmed
    book.create_sheet("Empty")
    staff = book.create_sheet("Staff")
    staff.append(["Asha"])
    source = tmp_path / "book.xlsx"
    book.save(source)

    model = _read(source, R.XLSX, make_request)

    assert [page.name for page in model.pages] == ["Sales", "Staff"]  # the empty sheet is skipped
    sales_table, staff_table = model.tables
    assert sales_table.rows == (("item", "qty"), ("Paper", "10"))
    assert sales_table.header_rows == 1
    assert staff_table.header_rows == 1


def test_keeps_empty_rows_between_data_when_reading_a_workbook(tmp_path: Path, make_request: MakeRequest) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet["A1"], sheet["A3"] = "top", "bottom"
    source = tmp_path / "gaps.xlsx"
    book.save(source)

    (table,) = _read(source, R.XLSX, make_request).tables

    assert table.rows == (("top",), ("",), ("bottom",))


def test_makes_one_unnamed_page_when_reading_a_csv(tmp_path: Path, make_request: MakeRequest) -> None:
    source = tmp_path / "notes.csv"
    source.write_text("1958,1959\n340,360\n", encoding="utf-8")

    model = _read(source, R.CSV, make_request)

    (page,) = model.pages
    assert page.name is None
    assert page.tables[0].header_rows == 0  # all numbers: no header can be proven


def test_writes_an_empty_model_when_the_csv_has_no_data(tmp_path: Path, make_request: MakeRequest) -> None:
    source = tmp_path / "blank.csv"
    source.write_text("\n\n", encoding="utf-8")

    assert _read(source, R.CSV, make_request).pages == ()
