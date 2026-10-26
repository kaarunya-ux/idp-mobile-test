"""PowerPoint -> model: a page per slide, text roles from bullets, tables, chart data, reading order."""

from collections.abc import Callable
from pathlib import Path

import pptx
import pytest
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.presentation import Presentation
from pptx.slide import Slide
from pptx.util import Inches

from conversion.conversion.domain.document_model import Alignment, DocumentModel, Table, TextBlock, TextKind
from conversion.conversion.domain.exceptions import UnreadableDocumentError
from conversion.conversion.domain.formats import Representation
from conversion.conversion.domain.ports import StepRequest
from conversion.conversion.domain.route import Converter
from conversion.conversion.infrastructure.converters.pptx_to_model import PythonPptxReader
from conversion.conversion.infrastructure.model_file import read_model

R = Representation
MakeRequest = Callable[..., StepRequest]
TITLE_AND_CONTENT, BLANK = 1, 6
BOX = (Inches(1), Inches(1), Inches(4), Inches(1))


def _read(deck: Presentation, tmp_path: Path, make_request: MakeRequest) -> DocumentModel:
    source = tmp_path / "deck.pptx"
    deck.save(str(source))
    request = make_request(Converter.PPTX_TO_MODEL, R.PPTX, R.DOCUMENT_MODEL, [source])
    (output,) = PythonPptxReader().convert(request).outputs
    return read_model(output)


def _blank_slide(deck: Presentation) -> Slide:
    return deck.slides.add_slide(deck.slide_layouts[BLANK])


def test_names_each_page_after_its_slide_title_or_number(tmp_path: Path, make_request: MakeRequest) -> None:
    deck = pptx.Presentation()
    titled = deck.slides.add_slide(deck.slide_layouts[TITLE_AND_CONTENT])
    titled.shapes.title.text = "Quarterly\vresults"  # a line break inside the title
    _blank_slide(deck)

    model = _read(deck, tmp_path, make_request)

    assert [page.name for page in model.pages] == ["Quarterly results", "Slide 2"]
    assert model.pages[0].blocks == ()  # the title names the page and isn't repeated as text


def test_reads_bullets_with_their_depth_when_the_layout_gives_them(tmp_path: Path, make_request: MakeRequest) -> None:
    deck = pptx.Presentation()
    body = deck.slides.add_slide(deck.slide_layouts[TITLE_AND_CONTENT]).placeholders[1].text_frame
    body.text = "Revenue up"
    nested = body.add_paragraph()
    nested.text, nested.level = "Mostly exports", 1
    plain = body.add_paragraph()
    plain.text = "Not a bullet"
    plain._p.get_or_add_pPr().append(plain._p.get_or_add_pPr().makeelement(qn("a:buNone"), {}))

    blocks = _read(deck, tmp_path, make_request).pages[0].blocks

    assert [(b.kind, b.level, b.text) for b in blocks if isinstance(b, TextBlock)] == [
        (TextKind.LIST_ITEM, 1, "Revenue up"),
        (TextKind.LIST_ITEM, 2, "Mostly exports"),
        (TextKind.PARAGRAPH, 1, "Not a bullet"),
    ]


def test_reads_a_bullet_written_in_a_text_box_and_keeps_alignment_and_bold(
    tmp_path: Path, make_request: MakeRequest
) -> None:
    deck = pptx.Presentation()
    frame = _blank_slide(deck).shapes.add_textbox(*BOX).text_frame
    frame.text = "Centred"
    frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    frame.paragraphs[0].runs[0].font.bold = True
    bullet = frame.add_paragraph()
    bullet.text = "Ticked"
    bullet._p.get_or_add_pPr().append(bullet._p.get_or_add_pPr().makeelement(qn("a:buChar"), {"char": "*"}))
    frame.add_paragraph().text = "  "  # empty paragraphs are skipped

    blocks = _read(deck, tmp_path, make_request).pages[0].blocks

    assert blocks == (
        TextBlock(text="Centred", alignment=Alignment.CENTER, bold=True),
        TextBlock(text="Ticked", kind=TextKind.LIST_ITEM),
    )


def test_reads_tables_with_their_header_row_and_merged_cells_once(tmp_path: Path, make_request: MakeRequest) -> None:
    deck = pptx.Presentation()
    grid = _blank_slide(deck).shapes.add_table(3, 2, *BOX).table  # "Header Row" is on by default
    grid.cell(0, 0).merge(grid.cell(0, 1))
    grid.cell(0, 0).text = "Sales"
    for row, (item, qty) in enumerate([("Paper", "10"), ("Ink", "2")], start=1):
        grid.cell(row, 0).text, grid.cell(row, 1).text = item, qty

    (table,) = _read(deck, tmp_path, make_request).tables

    assert table == Table(rows=(("Sales", ""), ("Paper", "10"), ("Ink", "2")), header_rows=1)


@pytest.mark.parametrize(("header_option", "rows", "expected"), [(False, 2, 0), (True, 1, 0)])
def test_reads_no_header_when_the_option_is_off_or_there_is_no_data_row(
    tmp_path: Path, make_request: MakeRequest, header_option: bool, rows: int, expected: int
) -> None:
    deck = pptx.Presentation()
    grid = _blank_slide(deck).shapes.add_table(rows, 1, *BOX).table
    grid.first_row = header_option
    grid.cell(0, 0).text = "x"

    (table,) = _read(deck, tmp_path, make_request).tables

    assert table.header_rows == expected


def test_skips_a_table_when_every_cell_is_empty(tmp_path: Path, make_request: MakeRequest) -> None:
    deck = pptx.Presentation()
    _blank_slide(deck).shapes.add_table(2, 2, *BOX)

    assert _read(deck, tmp_path, make_request).tables == []


def test_turns_a_charts_own_data_into_a_table_under_a_caption(tmp_path: Path, make_request: MakeRequest) -> None:
    deck = pptx.Presentation()
    data = CategoryChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("North", (120.0, 8.200000000000001))
    data.add_series("", (4.25, None))  # an unnamed series, with a missing value
    chart = _blank_slide(deck).shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, *BOX, data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Sales"

    caption, table = _read(deck, tmp_path, make_request).pages[0].blocks

    assert caption == TextBlock(text="Chart: Sales", kind=TextKind.CAPTION)
    assert table == Table(
        rows=(("Category", "North", "Series 2"), ("Q1", "120", "4.25"), ("Q2", "8.2", "")), header_rows=1
    )


def test_keeps_only_a_caption_when_a_chart_has_no_categories(tmp_path: Path, make_request: MakeRequest) -> None:
    deck = pptx.Presentation()
    data = XyChartData()
    data.add_series("Points").add_data_point(1.0, 2.0)
    _blank_slide(deck).shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, *BOX, data)

    blocks = _read(deck, tmp_path, make_request).pages[0].blocks

    assert blocks == (TextBlock(text="Chart", kind=TextKind.CAPTION),)


def test_reads_shapes_top_to_bottom_then_left_to_right_and_inside_groups(
    tmp_path: Path, make_request: MakeRequest, make_image: Callable[..., Path]
) -> None:
    deck = pptx.Presentation()
    shapes = _blank_slide(deck).shapes
    shapes.add_textbox(Inches(1), Inches(5), Inches(2), Inches(1)).text = "Bottom"
    shapes.add_textbox(Inches(5), Inches(1), Inches(2), Inches(1)).text = "Top right"
    shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = "Top left"
    group = shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(1)).text = "In a group"
    shapes.add_picture(str(make_image("photo.png")), Inches(1), Inches(4))  # pictures aren't read yet

    blocks = _read(deck, tmp_path, make_request).pages[0].blocks

    assert [b.text for b in blocks if isinstance(b, TextBlock)] == ["Top left", "Top right", "In a group", "Bottom"]


def test_reports_cnv_422_file_001_when_the_file_is_not_a_deck(tmp_path: Path, make_request: MakeRequest) -> None:
    fake = tmp_path / "fake.pptx"
    fake.write_bytes(b"not a zip at all")

    with pytest.raises(UnreadableDocumentError):
        PythonPptxReader().convert(make_request(Converter.PPTX_TO_MODEL, R.PPTX, R.DOCUMENT_MODEL, [fake]))
