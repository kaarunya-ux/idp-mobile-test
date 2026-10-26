"""Images -> PowerPoint: one slide per image, embedded unchanged, fitted and centred, upright."""

from collections.abc import Callable
from pathlib import Path

import pptx
import pytest
from pptx.shapes.picture import Picture

from conversion.conversion.domain.exceptions import LimitExceededError, UnreadableImageError
from conversion.conversion.domain.formats import Representation
from conversion.conversion.domain.ports import StepRequest
from conversion.conversion.domain.route import Converter
from conversion.conversion.infrastructure.converters.pptx_from_images import SLIDE_WIDTH_EMU, PptxImageSlides

R = Representation
MakeRequest = Callable[..., StepRequest]
MakeImage = Callable[..., Path]
MAX_PIXELS = 10_000


def _convert(
    make_request: MakeRequest, inputs: list[Path], max_pixels: int = MAX_PIXELS
) -> pptx.presentation.Presentation:
    request = make_request(Converter.PPTX_FROM_IMAGES, R.PNG, R.PPTX, inputs)
    result = PptxImageSlides(max_pixels=max_pixels).convert(request)
    (output,) = result.outputs
    assert output.name == f"{inputs[0].stem}.pptx"
    assert result.page_count == len(inputs)
    return pptx.Presentation(str(output))


def _picture(slide: pptx.slide.Slide) -> Picture:
    (picture,) = [shape for shape in slide.shapes if isinstance(shape, Picture)]
    return picture


def test_makes_one_slide_per_image_shaped_like_the_first(make_request: MakeRequest, make_image: MakeImage) -> None:
    portrait = make_image("page.png", size=(50, 100))

    deck = _convert(make_request, [portrait, make_image("wide.jpg", size=(100, 25))])

    assert (deck.slide_width, deck.slide_height) == (SLIDE_WIDTH_EMU, SLIDE_WIDTH_EMU * 2)
    assert len(deck.slides) == 2
    wide = _picture(deck.slides[1])
    assert wide.width == SLIDE_WIDTH_EMU  # fitted to the slide's width, not cropped
    assert wide.top == (SLIDE_WIDTH_EMU * 2 - wide.height) // 2  # centred


def test_embeds_the_image_bytes_unchanged(make_request: MakeRequest, make_image: MakeImage) -> None:
    photo = make_image("photo.jpg")

    deck = _convert(make_request, [photo])

    assert _picture(deck.slides[0]).image.blob == photo.read_bytes()


def test_turns_a_photo_upright_when_its_exif_says_to_rotate(make_request: MakeRequest, make_image: MakeImage) -> None:
    sideways = make_image("phone.jpg", size=(40, 20), orientation=6)  # shown rotated: 20 wide, 40 tall

    deck = _convert(make_request, [sideways])

    picture = _picture(deck.slides[0])
    assert picture.image.content_type == "image/png"
    assert picture.image.size == (20, 40)
    assert deck.slide_height == SLIDE_WIDTH_EMU * 2


def test_keeps_the_slide_within_powerpoints_size_limits(make_request: MakeRequest, make_image: MakeImage) -> None:
    banner = make_image("banner.png", size=(1000, 2))

    deck = _convert(make_request, [banner], max_pixels=1_000_000)

    assert deck.slide_height == 914_400  # 1 inch, PowerPoint's minimum


@pytest.mark.parametrize(
    ("content", "error"),
    [(b"not an image", UnreadableImageError), (None, LimitExceededError)],
)
def test_reports_a_domain_error_when_an_image_is_damaged_or_too_big(
    tmp_path: Path, make_request: MakeRequest, make_image: MakeImage, content: bytes | None, error: type[Exception]
) -> None:
    if content is None:
        source = make_image("huge.png", size=(200, 200))  # 40,000 pixels, over the 10,000 limit
    else:
        source = tmp_path / "broken.png"
        source.write_bytes(content)

    with pytest.raises(error):
        _convert(make_request, [source])
