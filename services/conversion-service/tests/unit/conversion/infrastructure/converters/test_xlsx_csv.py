"""CSV <-> XLSX: typed cells, IDs kept as text, encodings and delimiters detected, formulas never injected."""

from collections.abc import Callable
from datetime import datetime, time
from pathlib import Path

import openpyxl
import pytest

from conversion.conversion.domain.exceptions import LimitExceededError, SheetNotFoundError, UnreadableDocumentError
from conversion.conversion.domain.formats import Representation
from conversion.conversion.domain.options import ConversionOptions
from conversion.conversion.domain.ports import StepRequest
from conversion.conversion.domain.route import Converter
from conversion.conversion.infrastructure.converters import spreadsheets
from conversion.conversion.infrastructure.converters.xlsx_csv import OpenpyxlCsvConverter

R = Representation
MakeRequest = Callable[..., StepRequest]
converter = OpenpyxlCsvConverter()


def _to_xlsx(source: Path, make_request: MakeRequest) -> openpyxl.worksheet.worksheet.Worksheet:
    (output,) = converter.convert(make_request(Converter.XLSX_CSV, R.CSV, R.XLSX, [source])).outputs
    sheet = openpyxl.load_workbook(output).active
    assert sheet is not None
    return sheet


def _to_csv(source: Path, make_request: MakeRequest, sheet: str | None = None) -> list[str]:
    request = make_request(Converter.XLSX_CSV, R.XLSX, R.CSV, [source], ConversionOptions(sheet=sheet))
    (output,) = converter.convert(request).outputs
    return output.read_text(encoding="utf-8").splitlines()


def _csv(tmp_path: Path, text: str, *, encoding: str = "utf-8", name: str = "data.csv") -> Path:
    path = tmp_path / name
    path.write_bytes(text.encode(encoding))
    return path


def test_types_numbers_and_keeps_ids_as_text_when_converting_csv(tmp_path: Path, make_request: MakeRequest) -> None:
    source = _csv(tmp_path, 'item,qty,price,pin\nPaper,10,"1,250.50",0471\nCard,1,9.99,4111111111111111\n')

    sheet = _to_xlsx(source, make_request)

    assert [c.value for c in sheet[1]] == ["item", "qty", "price", "pin"]
    assert sheet["A1"].font.bold
    assert (sheet["C2"].value, sheet["C2"].number_format) == (1250.5, "#,##0.00")
    assert (sheet["D2"].value, sheet["D2"].number_format) == ("0471", "@")  # a PIN keeps its leading zero
    assert sheet["D3"].value == "4111111111111111"  # a card number is never rounded by Excel


def test_keeps_formula_like_text_as_text_when_a_cell_starts_with_equals(
    tmp_path: Path, make_request: MakeRequest
) -> None:
    source = _csv(tmp_path, 'name,note\nAsha,"=HYPERLINK(""http://evil"")"\n')

    sheet = _to_xlsx(source, make_request)

    assert sheet["B2"].data_type == "s"
    assert sheet["B2"].value == '=HYPERLINK("http://evil")'


@pytest.mark.parametrize(
    ("text", "encoding"),
    [("name;city\nRené;Café\n", "cp1252"), ("name;city\nRené;Café\n", "utf-16"), ("name\tcity\nRené\tCafé\n", "utf-8")],
)
def test_detects_the_encoding_and_delimiter_when_reading_csv(
    tmp_path: Path, make_request: MakeRequest, text: str, encoding: str
) -> None:
    sheet = _to_xlsx(_csv(tmp_path, text, encoding=encoding), make_request)

    assert [c.value for c in sheet[2]] == ["René", "Café"]


def test_keeps_decimal_comma_numbers_as_text_when_the_file_uses_them(tmp_path: Path, make_request: MakeRequest) -> None:
    sheet = _to_xlsx(_csv(tmp_path, "item;price\nPaper;1.250,50\nPens;96,00\n"), make_request)

    assert (sheet["B2"].value, sheet["B2"].data_type) == ("1.250,50", "s")


def test_writes_an_empty_sheet_when_the_csv_has_no_data(tmp_path: Path, make_request: MakeRequest) -> None:
    sheet = _to_xlsx(_csv(tmp_path, "\n\n"), make_request)

    assert sheet.max_row == 1
    assert sheet["A1"].value is None


def test_reports_cnv_422_file_001_when_the_csv_is_really_a_binary_file(
    tmp_path: Path, make_request: MakeRequest
) -> None:
    source = tmp_path / "photo.csv"
    source.write_bytes(b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR")

    with pytest.raises(UnreadableDocumentError) as caught:
        _to_xlsx(source, make_request)

    assert caught.value.code == "CNV-422-FILE-001"


def test_refuses_when_the_table_is_bigger_than_a_spreadsheet(
    tmp_path: Path, make_request: MakeRequest, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(spreadsheets, "EXCEL_MAX_ROWS", 2)

    with pytest.raises(LimitExceededError):
        _to_xlsx(_csv(tmp_path, "a\n1\n2\n3\n"), make_request)


@pytest.fixture
def workbook(tmp_path: Path) -> Path:
    book = openpyxl.Workbook()
    first = book.active
    assert first is not None
    first.title = "Sales"
    first.append(["item", "qty", "price", "when", "paid"])
    first.append(["Paper", 10, 0.1, datetime(2026, 10, 21), True])  # noqa: DTZ001 - a spreadsheet date has no zone
    first.append(["Pens", 3.0, 1e-05, datetime(2026, 10, 21, 9, 30), False])  # noqa: DTZ001 - same
    first["A6"] = "Ink"  # rows 4-5 stay empty: a gap in the middle is kept
    first["C6"] = time(9, 30)  # a time of day; B6 stays empty
    first["A9"].number_format = "0.00"  # a styled but empty cell: trailing empty rows are not exported
    other = book.create_sheet("Staff")
    other.append(["name"])
    other.append(["Asha"])
    path = tmp_path / "book.xlsx"
    book.save(path)
    return path


def test_exports_the_first_sheet_as_written_values_when_no_sheet_is_named(
    workbook: Path, make_request: MakeRequest
) -> None:
    lines = _to_csv(workbook, make_request)

    assert lines == [
        "item,qty,price,when,paid",
        "Paper,10,0.1,2026-10-21,TRUE",
        "Pens,3,0.00001,2026-10-21 09:30:00,FALSE",
        "",
        "",
        "Ink,,09:30:00,,",
    ]


def test_exports_the_named_sheet_when_the_sheet_option_is_set(workbook: Path, make_request: MakeRequest) -> None:
    assert _to_csv(workbook, make_request, sheet="Staff") == ["name", "Asha"]


def test_reports_cnv_422_options_001_when_the_named_sheet_does_not_exist(
    workbook: Path, make_request: MakeRequest
) -> None:
    with pytest.raises(SheetNotFoundError) as caught:
        _to_csv(workbook, make_request, sheet="Budget")

    assert caught.value.code == "CNV-422-OPTIONS-001"


def test_reports_cnv_422_file_001_when_the_workbook_is_damaged(tmp_path: Path, make_request: MakeRequest) -> None:
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(b"PK\x03\x04 not really a workbook")

    with pytest.raises(UnreadableDocumentError):
        _to_csv(broken, make_request)


def test_round_trips_text_and_ids_exactly_when_going_csv_to_xlsx_and_back(
    tmp_path: Path, make_request: MakeRequest
) -> None:
    original = "name,pin,city\nAsha,0471,Pune\nRavi,0560,Delhi\n"
    sheet_file = converter.convert(make_request(Converter.XLSX_CSV, R.CSV, R.XLSX, [_csv(tmp_path, original)])).outputs[
        0
    ]

    assert _to_csv(sheet_file, make_request) == original.splitlines()


def test_reports_cnv_422_file_001_when_a_csv_field_is_too_large_to_parse(
    tmp_path: Path, make_request: MakeRequest
) -> None:
    source = _csv(tmp_path, 'note\n"' + "x" * 200_000 + '"\n')  # past Python's csv field limit (131,072)

    with pytest.raises(UnreadableDocumentError):
        _to_xlsx(source, make_request)


def test_refuses_a_target_this_converter_does_not_write(tmp_path: Path, make_request: MakeRequest) -> None:
    request = make_request(Converter.XLSX_CSV, R.CSV, R.PDF, [_csv(tmp_path, "a\n1\n")])

    with pytest.raises(ValueError, match="can't write PDF"):
        converter.convert(request)


def test_still_reads_the_file_when_bytes_fit_no_common_encoding(tmp_path: Path, make_request: MakeRequest) -> None:
    source = tmp_path / "odd.csv"
    source.write_bytes(b"code\nA\x81B\n")  # 0x81 is invalid in UTF-8 and undefined in Windows-1252

    sheet = _to_xlsx(source, make_request)

    assert sheet["A2"].value == "A\x81B"  # read as Latin-1, never rejected
