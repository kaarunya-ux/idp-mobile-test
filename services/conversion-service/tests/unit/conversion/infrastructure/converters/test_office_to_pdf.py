"""Office -> PDF with LibreOffice: pages counted, Hindi rendered, damaged files reported, runaway runs stopped."""

import shutil
import time
from collections.abc import Callable
from datetime import timedelta
from pathlib import Path

import docx
import pikepdf
import pptx
import pytest

from conversion.conversion.domain.exceptions import ConversionTimeoutError, UnreadableDocumentError
from conversion.conversion.domain.formats import Representation
from conversion.conversion.domain.ports import StepRequest
from conversion.conversion.domain.route import Converter
from conversion.conversion.infrastructure.converters.office_to_pdf import (
    LibreOfficeNotInstalledError,
    LibreOfficePdfConverter,
)

R = Representation
MakeRequest = Callable[..., StepRequest]
needs_libreoffice = pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice is not installed")
converter = LibreOfficePdfConverter(binary="soffice", timeout=timedelta(seconds=60))


def _pdf_fonts(path: Path) -> set[str]:
    with pikepdf.open(path) as pdf:
        return {str(font.get("/BaseFont")) for page in pdf.pages for font in page.Resources.get("/Font", {}).values()}


@needs_libreoffice
def test_converts_a_word_file_with_hindi_and_embeds_a_devanagari_font(
    tmp_path: Path, make_request: MakeRequest
) -> None:
    document = docx.Document()
    document.add_paragraph("नमस्ते, यह हिन्दी में एक वाक्य है।")
    document.add_page_break()
    document.add_paragraph("Second page")
    source = tmp_path / "letter.docx"
    document.save(source)

    result = converter.convert(make_request(Converter.OFFICE_TO_PDF, R.DOCX, R.PDF, [source]))

    (output,) = result.outputs
    assert output.name == "letter.pdf"
    assert result.page_count == 2
    assert any("Devanagari" in font for font in _pdf_fonts(output))  # IDP-03 ch. 3.4: Noto incl. Devanagari


@needs_libreoffice
def test_converts_a_presentation_to_one_page_per_slide(tmp_path: Path, make_request: MakeRequest) -> None:
    deck = pptx.Presentation()
    for title in ("Q3", "Q4", "Plan"):
        deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = title
    source = tmp_path / "deck.pptx"
    deck.save(source)

    result = converter.convert(make_request(Converter.OFFICE_TO_PDF, R.PPTX, R.PDF, [source]))

    assert result.page_count == 3


@needs_libreoffice
def test_reports_cnv_422_file_001_when_the_document_is_damaged(tmp_path: Path, make_request: MakeRequest) -> None:
    source = tmp_path / "broken.docx"
    source.write_bytes(b"PK\x03\x04 this is not a real document")

    with pytest.raises(UnreadableDocumentError) as caught:
        converter.convert(make_request(Converter.OFFICE_TO_PDF, R.DOCX, R.PDF, [source]))

    assert caught.value.code == "CNV-422-FILE-001"


def test_stops_libreoffice_and_reports_cnv_504_timeout_001_when_it_runs_too_long(
    tmp_path: Path, make_request: MakeRequest
) -> None:
    hanging = tmp_path / "fake-soffice"
    hanging.write_text("#!/bin/sh\nsleep 30\n", encoding="utf-8")
    hanging.chmod(0o755)
    source = tmp_path / "letter.docx"
    source.write_bytes(b"x")
    slow = LibreOfficePdfConverter(binary=str(hanging), timeout=timedelta(seconds=0.5))

    started = time.monotonic()
    with pytest.raises(ConversionTimeoutError) as caught:
        slow.convert(make_request(Converter.OFFICE_TO_PDF, R.DOCX, R.PDF, [source]))

    assert caught.value.code == "CNV-504-TIMEOUT-001"
    assert time.monotonic() - started < 10  # killed, not waited out


def test_raises_a_deployment_error_when_libreoffice_is_missing(tmp_path: Path, make_request: MakeRequest) -> None:
    source = tmp_path / "letter.docx"
    source.write_bytes(b"x")
    missing = LibreOfficePdfConverter(binary="no-such-office-binary", timeout=timedelta(seconds=5))

    with pytest.raises(LibreOfficeNotInstalledError):
        missing.convert(make_request(Converter.OFFICE_TO_PDF, R.DOCX, R.PDF, [source]))
