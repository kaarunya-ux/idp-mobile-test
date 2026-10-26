"""Model -> Word, XLSX and CSV: styles, header rows, typed sheets, table choice and size limits."""

from collections.abc import Callable
from pathlib import Path

import docx
import openpyxl
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import RGBColor

from conversion.conversion.domain.document_model import (
    Alignment,
    DocumentModel,
    Page,
    Table,
    TextBlock,
    TextKind,
)
from conversion.conversion.domain.exceptions import LimitExceededError, NoTableFoundError, SheetNotFoundError
from conversion.conversion.domain.formats import Representation
from conversion.conversion.domain.options import ConversionOptions
from conversion.conversion.domain.ports import ConverterPort, StepRequest
from conversion.conversion.domain.route import Converter
from conversion.conversion.infrastructure.converters.model_to_docx import PythonDocxWriter
from conversion.conversion.infrastructure.converters.model_to_sheets import CsvModelWriter, OpenpyxlModelWriter
from conversion.conversion.infrastructure.model_file import write_model

R = Representation
MakeRequest = Callable[..., StepRequest]
SALES = Table(rows=(("item", "qty", "price"), ("Paper", "10", "1,250.50")), header_rows=1)
STAFF = Table(rows=(("name",), ("Asha",)), header_rows=1)


def _model_file(tmp_path: Path, model: DocumentModel) -> Path:
    path = tmp_path / "report.model.json"
    write_model(model, path)
    return path


def _write(
    converter: ConverterPort, kind: Converter, target: R, model_file: Path, make_request: MakeRequest, **options: str
) -> Path:
    request = make_request(kind, R.DOCUMENT_MODEL, target, [model_file], ConversionOptions(**options))
    (output,) = converter.convert(request).outputs
    return output


def _docx(tmp_path: Path, model: DocumentModel, make_request: MakeRequest, max_rows: int = 1000) -> Path:
    writer = PythonDocxWriter(max_table_rows=max_rows)
    return _write(writer, Converter.MODEL_TO_DOCX, R.DOCX, _model_file(tmp_path, model), make_request)


def test_writes_word_styles_for_each_role_when_making_a_document(tmp_path: Path, make_request: MakeRequest) -> None:
    model = DocumentModel(
        pages=(
            Page(
                name="Sales",
                blocks=(
                    TextBlock(text="Summary", kind=TextKind.HEADING, level=2),
                    TextBlock(text="Point", kind=TextKind.LIST_ITEM, level=5),
                    TextBlock(text="Figure 1", kind=TextKind.CAPTION),
                    TextBlock(text="Centred", alignment=Alignment.CENTER, bold=True),
                ),
            ),
        )
    )

    output = _docx(tmp_path, model, make_request)

    assert output.name == "report.docx"
    paragraphs = docx.Document(str(output)).paragraphs
    assert [(p.style.name, p.text) for p in paragraphs if p.style is not None] == [
        ("Heading 1", "Sales"),  # the page name
        ("Heading 2", "Summary"),
        ("List Bullet 3", "Point"),  # deeper levels use the deepest list style Word's template has
        ("Caption", "Figure 1"),
        ("Normal", "Centred"),
    ]
    assert paragraphs[-1].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert all(run.bold for run in paragraphs[-1].runs)


def test_makes_headings_black_when_writing_a_document(tmp_path: Path, make_request: MakeRequest) -> None:
    model = DocumentModel(pages=(Page(blocks=(TextBlock(text="Title", kind=TextKind.HEADING),)),))

    document = docx.Document(str(_docx(tmp_path, model, make_request)))

    assert document.styles["Heading 1"].font.color.rgb == RGBColor(0, 0, 0)


def test_writes_a_grid_whose_header_repeats_on_each_page(tmp_path: Path, make_request: MakeRequest) -> None:
    model = DocumentModel(pages=(Page(blocks=(SALES,)),))

    (table,) = docx.Document(str(_docx(tmp_path, model, make_request))).tables

    assert [[cell.text for cell in row.cells] for row in table.rows] == [list(row) for row in SALES.rows]
    header = table.rows[0]
    assert header._tr.trPr is not None
    assert header._tr.trPr.find(qn("w:tblHeader")) is not None
    assert all(run.bold for run in header.cells[0].paragraphs[0].runs)
    assert table.style is not None
    assert table.style.name == "Table Grid"


def test_refuses_a_table_too_long_for_a_document(tmp_path: Path, make_request: MakeRequest) -> None:
    long_table = Table(rows=tuple((str(n),) for n in range(201)))

    with pytest.raises(LimitExceededError, match="XLSX or CSV"):
        _docx(tmp_path, DocumentModel(pages=(Page(blocks=(long_table,)),)), make_request, max_rows=200)


def test_makes_one_typed_sheet_per_table_named_after_its_page(tmp_path: Path, make_request: MakeRequest) -> None:
    model = DocumentModel(
        pages=(
            Page(name="Sales/2026?", blocks=(SALES,)),
            Page(name="Two tables", blocks=(STAFF, STAFF)),
            Page(blocks=(TextBlock(text="text only"),)),
        )
    )

    output = _write(OpenpyxlModelWriter(), Converter.MODEL_TO_XLSX, R.XLSX, _model_file(tmp_path, model), make_request)

    book = openpyxl.load_workbook(output)
    assert book.sheetnames == ["Sales 2026", "Table 2", "Table 3"]  # forbidden characters removed; numbered otherwise
    sales = book["Sales 2026"]
    assert (sales["C2"].value, sales["C2"].number_format) == (1250.5, "#,##0.00")
    assert sales["A1"].font.bold


def test_reports_cnv_422_table_001_when_asked_for_a_sheet_from_a_document_without_tables(
    tmp_path: Path, make_request: MakeRequest
) -> None:
    model = DocumentModel(pages=(Page(blocks=(TextBlock(text="Just words"),)),))

    with pytest.raises(NoTableFoundError) as caught:
        _write(OpenpyxlModelWriter(), Converter.MODEL_TO_XLSX, R.XLSX, _model_file(tmp_path, model), make_request)

    assert caught.value.code == "CNV-422-TABLE-001"


def test_exports_the_first_table_when_no_sheet_is_named(tmp_path: Path, make_request: MakeRequest) -> None:
    model = DocumentModel(pages=(Page(name="Sales", blocks=(SALES,)), Page(name="Staff", blocks=(STAFF,))))

    output = _write(CsvModelWriter(), Converter.MODEL_TO_CSV, R.CSV, _model_file(tmp_path, model), make_request)

    assert output.read_text(encoding="utf-8").splitlines() == ["item,qty,price", 'Paper,10,"1,250.50"']


def test_exports_the_named_pages_table_when_the_sheet_option_is_set(tmp_path: Path, make_request: MakeRequest) -> None:
    model = DocumentModel(pages=(Page(name="Sales", blocks=(SALES,)), Page(name="Staff", blocks=(STAFF,))))
    model_file = _model_file(tmp_path, model)

    output = _write(CsvModelWriter(), Converter.MODEL_TO_CSV, R.CSV, model_file, make_request, sheet="Staff")

    assert output.read_text(encoding="utf-8").splitlines() == ["name", "Asha"]


@pytest.mark.parametrize(("sheet", "error"), [("Budget", SheetNotFoundError), ("Notes", NoTableFoundError)])
def test_reports_a_clear_error_when_the_named_page_is_missing_or_has_no_table(
    tmp_path: Path, make_request: MakeRequest, sheet: str, error: type[Exception]
) -> None:
    model = DocumentModel(
        pages=(Page(name="Sales", blocks=(SALES,)), Page(name="Notes", blocks=(TextBlock(text="x"),)))
    )

    with pytest.raises(error):
        _write(CsvModelWriter(), Converter.MODEL_TO_CSV, R.CSV, _model_file(tmp_path, model), make_request, sheet=sheet)
