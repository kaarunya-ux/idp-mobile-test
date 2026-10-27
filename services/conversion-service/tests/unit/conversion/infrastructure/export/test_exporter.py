"""Exports of extraction results: each format's layout, masked values kept masked, formulas never run."""

import csv
import json
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

import docx
import openpyxl
import pytest

from conversion.conversion.domain.exceptions import InvalidOptionError
from conversion.conversion.domain.formats import Format
from conversion.conversion.domain.results import ExportLayout, ExportOptions, parse_result
from conversion.conversion.infrastructure.export.exporter import XML_NAMESPACE, ResultExporter

exporter = ResultExporter()


def _result(name: str, fields: list[dict[str, Any]], tables: list[dict[str, Any]] | None = None):
    data = {
        "schemaVersion": 1,
        "documentType": {"code": "PAN_CARD", "confidence": 0.97},
        "pages": [{"index": 0, "imageKey": "acc/doc/page-0.png"}],
        "fields": fields,
        "tables": tables or [],
        "validation": {"fieldsTotal": len(fields), "accepted": 1, "needsReview": 1, "failedRules": ["dob-date"]},
    }
    return parse_result(data, name=name)


PAN = {
    "key": "pan",
    "label": "PAN",
    "displayValue": "XXXXX1234F",
    "confidence": 0.99,
    "source": "OCR_AGREE",
    "masked": True,
}
NAME = {"key": "name", "label": "Name", "displayValue": '=HYPERLINK("x")', "confidence": 0.8, "status": "NEEDS_REVIEW"}
DOB = {"key": "dob", "label": "Date of birth", "value": "1990-01-01", "confidence": 0.9}
RESULTS = [_result("doc-1", [PAN, NAME]), _result("doc-2", [DOB, PAN])]


def _options(layout: ExportLayout | None = None, *, confidence: bool = False) -> ExportOptions:
    return ExportOptions(layout=layout, include_confidence=confidence)


def test_writes_one_array_or_one_file_per_result_when_exporting_json(tmp_path: Path) -> None:
    (array,) = exporter.export(RESULTS, Format.JSON, _options(), tmp_path / "a", {})
    per_file = exporter.export(RESULTS, Format.JSON, _options(ExportLayout.ONE_FILE_PER_DOCUMENT), tmp_path / "b", {})

    assert [r["fields"][0]["key"] for r in json.loads(array.read_text())] == ["pan", "dob"]
    assert [p.name for p in per_file] == ["results-001.json", "results-002.json"]


def test_writes_a_row_per_document_with_the_union_of_fields_when_exporting_csv(tmp_path: Path) -> None:
    (path,) = exporter.export(RESULTS, Format.CSV, _options(confidence=True), tmp_path, {})

    rows = list(csv.reader(path.open(encoding="utf-8-sig")))
    assert rows[0] == [
        "document",
        "documentType",
        "pan",
        "pan_confidence",
        "name",
        "name_confidence",
        "dob",
        "dob_confidence",
    ]
    assert rows[1] == ["doc-1", "PAN_CARD", "XXXXX1234F", "0.99", '\'=HYPERLINK("x")', "0.8", "", ""]
    assert rows[2][2] == "XXXXX1234F"  # masked in every export


def test_writes_a_documents_sheet_and_a_sheet_per_table_type_when_exporting_xlsx(tmp_path: Path) -> None:
    table = {
        "key": "line_items",
        "columns": [{"key": "item"}, {"key": "amount"}],
        "rows": [[{"value": "Paper"}, {"value": 12}]],
    }
    results = [_result("doc-1", [PAN], [table])]

    (path,) = exporter.export(results, Format.XLSX, _options(), tmp_path, {})

    workbook = openpyxl.load_workbook(path)
    assert workbook.sheetnames == ["Documents", "line_items"]
    assert [c.value for c in workbook["line_items"][2]] == ["doc-1", "Paper", "12"]


def test_writes_a_sheet_per_document_when_the_layout_asks_for_it(tmp_path: Path) -> None:
    (path,) = exporter.export(RESULTS, Format.XLSX, _options(ExportLayout.ONE_SHEET_PER_DOCUMENT), tmp_path, {})

    workbook = openpyxl.load_workbook(path)
    assert workbook.sheetnames == ["doc-1", "doc-2"]
    assert workbook["doc-1"]["B3"].value == '\'=HYPERLINK("x")'  # stored as text, never run


def test_writes_a_readable_report_with_the_page_image_when_exporting_docx(tmp_path: Path, make_image) -> None:
    thumbnail = make_image("page.png", size=(100, 60))

    (path,) = exporter.export(RESULTS, Format.DOCX, _options(), tmp_path, {"doc-1": thumbnail})

    report = docx.Document(str(path))
    text = "\n".join(p.text for p in report.paragraphs)
    assert "PAN_CARD: doc-1" in text
    assert "Failed checks: dob-date" in text
    assert report.tables[0].rows[1].cells[1].text == "XXXXX1234F"
    assert len(report.inline_shapes) == 1


def test_writes_an_element_per_field_in_the_idp_namespace_when_exporting_xml(tmp_path: Path) -> None:
    (path,) = exporter.export(RESULTS, Format.XML, _options(), tmp_path, {})

    # our own output
    root = ET.parse(path).getroot()  # noqa: S314
    fields = root.findall(f"{{{XML_NAMESPACE}}}document/{{{XML_NAMESPACE}}}field")
    assert root.tag == f"{{{XML_NAMESPACE}}}results"
    assert (fields[0].get("key"), fields[0].text, fields[0].get("masked")) == ("pan", "XXXXX1234F", "true")


@pytest.mark.parametrize(
    ("target", "layout"), [(Format.CSV, ExportLayout.ARRAY), (Format.JSON, ExportLayout.ONE_SHEET_PER_DOCUMENT)]
)
def test_reports_cnv_422_options_001_when_the_layout_does_not_fit_the_format(
    tmp_path: Path, target: Format, layout: ExportLayout
) -> None:
    with pytest.raises(InvalidOptionError):
        exporter.export(RESULTS, target, _options(layout), tmp_path, {})
