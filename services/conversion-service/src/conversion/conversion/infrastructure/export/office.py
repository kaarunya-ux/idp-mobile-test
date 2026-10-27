"""XLSX sheets and the DOCX report for ``EXPORT`` (openpyxl, python-docx; IDP-03 ch. 3.6)."""

import re
from collections.abc import Mapping, Sequence
from pathlib import Path

import docx
import openpyxl
from docx.document import Document as DocxDocument
from docx.shared import Cm
from openpyxl.worksheet.worksheet import Worksheet

from conversion.conversion.domain.results import ExportLayout, ExportOptions, ResultDocument, ResultTable
from conversion.conversion.infrastructure.export.rows import header_and_rows

_SHEET_NAME_LIMIT = 31  # Excel's own limit
_BAD_SHEET_CHARACTERS = re.compile(r"[\[\]:*?/\\]")
_THUMBNAIL_WIDTH = Cm(8)


def xlsx(results: Sequence[ResultDocument], options: ExportOptions, path: Path) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    if not isinstance(sheet, Worksheet):  # a new workbook always has one sheet
        raise TypeError("openpyxl made a workbook without a worksheet")
    names: set[str] = set()
    if options.layout is ExportLayout.ONE_SHEET_PER_DOCUMENT:
        workbook.remove(sheet)
        for result in results:
            _document_sheet(workbook.create_sheet(_sheet_name(result.name, names)), result, options)
    else:
        sheet.title = _sheet_name("Documents", names)
        for row in header_and_rows(results, include_confidence=options.include_confidence):
            sheet.append([_text_cell(cell) for cell in row])
        for key, tables in _tables_by_key(results).items():
            _table_sheet(workbook.create_sheet(_sheet_name(key, names)), tables)
    workbook.save(path)
    return path


def docx_report(results: Sequence[ResultDocument], thumbnails: Mapping[str, Path], path: Path) -> Path:
    """A readable report: per document, its first page, the fields, its tables and the validation summary."""
    report = docx.Document()
    report.add_heading("Extraction results", level=0)
    for number, result in enumerate(results):
        if number:
            report.add_page_break()  # type: ignore[no-untyped-call]  # python-docx: no hints on this method
        _document_section(report, result, thumbnails.get(result.name))
    report.save(str(path))
    return path


def _document_section(report: DocxDocument, result: ResultDocument, thumbnail: Path | None) -> None:
    report.add_heading(f"{result.document_type or 'Document'}: {result.name}", level=1)
    if thumbnail is not None:
        report.add_picture(str(thumbnail), width=_THUMBNAIL_WIDTH)
    _fields_table(report, result)
    for table in result.tables:
        _docx_table(report, table)
    report.add_paragraph(_summary(result))


def _fields_table(report: DocxDocument, result: ResultDocument) -> None:
    fields = report.add_table(rows=1, cols=4)
    fields.style = "Table Grid"
    for cell, title in zip(fields.rows[0].cells, ("Field", "Value", "Confidence", "Source"), strict=True):
        cell.text = title
    for field in result.fields:
        cells = fields.add_row().cells
        cells[0].text, cells[1].text = field.label, field.value or ""
        cells[2].text = f"{field.confidence:.0%}" if field.confidence is not None else ""
        cells[3].text = field.source or ""


def _summary(result: ResultDocument) -> str:
    summary = f"Accepted: {result.accepted}   Needs review: {result.needs_review}"
    if result.failed_rules:
        summary += f"   Failed checks: {', '.join(result.failed_rules)}"
    return summary


def _document_sheet(sheet: Worksheet, result: ResultDocument, options: ExportOptions) -> None:
    header = ["Field", "Value", "Status", "Source"] + (["Confidence"] if options.include_confidence else [])
    sheet.append(header)
    for field in result.fields:
        row: list[object] = [field.label, _text_cell(field.value or ""), field.status or "", field.source or ""]
        if options.include_confidence:
            row.append(field.confidence if field.confidence is not None else "")
        sheet.append(row)


def _tables_by_key(results: Sequence[ResultDocument]) -> dict[str, list[tuple[str, ResultTable]]]:
    """One sheet per table type, across documents (IDP-03 ch. 3.6)."""
    grouped: dict[str, list[tuple[str, ResultTable]]] = {}
    for result in results:
        for table in result.tables:
            grouped.setdefault(table.key, []).append((result.name, table))
    return grouped


def _table_sheet(sheet: Worksheet, tables: Sequence[tuple[str, ResultTable]]) -> None:
    columns = max((t.columns for _, t in tables), key=len)
    sheet.append(["document", *columns])
    for name, table in tables:
        for row in table.rows:
            sheet.append([name, *(_text_cell(cell) for cell in row)])


def _docx_table(report: DocxDocument, table: ResultTable) -> None:
    report.add_heading(table.key, level=2)
    grid = report.add_table(rows=1, cols=max(len(table.columns), 1))
    grid.style = "Table Grid"
    for cell, title in zip(grid.rows[0].cells, table.columns, strict=False):
        cell.text = title
    for row in table.rows:
        cells = grid.add_row().cells
        for cell, value in zip(cells, row, strict=False):
            cell.text = value


def _sheet_name(wanted: str, used: set[str]) -> str:
    base = _BAD_SHEET_CHARACTERS.sub("_", wanted)[:_SHEET_NAME_LIMIT] or "Sheet"
    name, number = base, 2
    while name.lower() in used:
        suffix = f" ({number})"
        name, number = base[: _SHEET_NAME_LIMIT - len(suffix)] + suffix, number + 1
    used.add(name.lower())
    return name


def _text_cell(value: object) -> object:
    """Strings that look like formulas are stored as text, so opening the file runs nothing (CSV/XLSX injection)."""
    if isinstance(value, str) and value.startswith("="):  # openpyxl makes formulas only of "=..."
        return f"'{value}"
    return value
