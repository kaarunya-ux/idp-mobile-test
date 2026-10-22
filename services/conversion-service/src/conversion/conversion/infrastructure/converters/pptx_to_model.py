"""PowerPoint -> document model with python-pptx (IDP-03 ch. 3.4: "PPTX -> DOCX/XLSX/CSV").

Ported from ``ppt_to_office.py`` (Karunya, Sumukh) without the AI step. Everything comes from what the deck
stores: each slide becomes a page named after its title; text, tables and the data behind charts become blocks
in reading order (top to bottom, then left to right). Pictures, SmartArt and speaker notes are not read yet.
"""

import zipfile
from collections.abc import Iterable
from pathlib import Path

import pptx
from pptx.chart.chart import Chart
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN
from pptx.exc import PackageNotFoundError
from pptx.oxml.ns import qn
from pptx.presentation import Presentation
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pptx.slide import Slide
from pptx.table import Table as SlideTable
from pptx.text.text import _Paragraph

from conversion.conversion.domain.document_model import (
    MAX_HEADING_LEVEL,
    Alignment,
    Block,
    DocumentModel,
    Page,
    Table,
    TextBlock,
    TextKind,
)
from conversion.conversion.domain.exceptions import UnreadableDocumentError
from conversion.conversion.domain.ports import StepRequest, StepResult
from conversion.conversion.domain.route import Converter
from conversion.conversion.infrastructure.model_file import model_path, write_model

_ROW_BAND_EMU = 228_600  # 0.25 inch: shapes whose tops are this close read as one row, then left to right
_BULLETED_PLACEHOLDERS = frozenset({PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT})  # bullets come from the layout
_ALIGNMENTS = {PP_ALIGN.CENTER: Alignment.CENTER, PP_ALIGN.RIGHT: Alignment.RIGHT}
_MIN_ROWS_FOR_A_HEADER = 2  # a header row needs at least one data row under it
_CHART_DIGITS = 12  # chart values are stored as doubles: 8.200000000000001 is shown as 8.2


class PythonPptxReader:
    """Writes ``<input>.model.json``: one page per slide, hidden slides included."""

    @property
    def converter(self) -> Converter:
        return Converter.PPTX_TO_MODEL

    def convert(self, request: StepRequest) -> StepResult:
        (source,) = request.inputs
        deck = _open(source)
        pages = tuple(_page(slide, number) for number, slide in enumerate(deck.slides, start=1))
        output = model_path(request.output_dir, source)
        write_model(DocumentModel(pages=pages), output)
        return StepResult(outputs=(output,), page_count=max(len(pages), 1))


def _open(source: Path) -> Presentation:
    try:
        return pptx.Presentation(str(source))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:  # not a PowerPoint file
        raise UnreadableDocumentError from exc


def _page(slide: Slide, number: int) -> Page:
    """The slide's title names the page; it isn't repeated as a block."""
    title_shape = slide.shapes.title
    title = _one_line(title_shape.text_frame.text) if title_shape is not None else ""
    title_id = title_shape.shape_id if title_shape is not None else None
    blocks = [block for shape in _reading_order(slide.shapes) if shape.shape_id != title_id for block in _blocks(shape)]
    return Page(name=title or f"Slide {number}", blocks=tuple(blocks))


def _reading_order(shapes: Iterable[BaseShape]) -> list[BaseShape]:
    """Shapes top to bottom in bands, then left to right; a group's shapes are read in place of the group."""
    ordered: list[BaseShape] = []
    for shape in sorted(shapes, key=_position):
        if isinstance(shape, GroupShape):
            ordered.extend(_reading_order(shape.shapes))
        else:
            ordered.append(shape)
    return ordered


def _position(shape: BaseShape) -> tuple[int, int]:
    return round((shape.top or 0) / _ROW_BAND_EMU), shape.left or 0


def _blocks(shape: BaseShape) -> list[Block]:
    if isinstance(shape, GraphicFrame):
        if shape.has_table:
            table = _table(shape.table)
            return [table] if table is not None else []
        return _chart(shape.chart) if shape.has_chart else []
    if isinstance(shape, Shape) and shape.has_text_frame:  # text boxes, placeholders and shapes with text
        return [block for paragraph in shape.text_frame.paragraphs if (block := _text(paragraph, shape))]
    return []  # pictures, lines, embedded objects


def _text(paragraph: _Paragraph, shape: Shape) -> TextBlock | None:
    text = _one_line(paragraph.text)
    if not text:
        return None
    runs = [run for run in paragraph.runs if run.text.strip()]
    bulleted = _is_bulleted(paragraph, shape)
    return TextBlock(
        text=text,
        kind=TextKind.LIST_ITEM if bulleted else TextKind.PARAGRAPH,
        level=min(paragraph.level + 1, MAX_HEADING_LEVEL) if bulleted else 1,
        alignment=_ALIGNMENTS.get(paragraph.alignment, Alignment.LEFT) if paragraph.alignment else Alignment.LEFT,
        bold=bool(runs) and all(run.font.bold for run in runs),
    )


def _is_bulleted(paragraph: _Paragraph, shape: Shape) -> bool:
    """A bullet or number written on the paragraph, else whatever a body placeholder gives it from the layout."""
    properties = paragraph._p.pPr  # bullets are only reachable through the XML element
    if properties is not None:
        if properties.find(qn("a:buNone")) is not None:
            return False
        if properties.find(qn("a:buChar")) is not None or properties.find(qn("a:buAutoNum")) is not None:
            return True
    return shape.is_placeholder and shape.placeholder_format.type in _BULLETED_PLACEHOLDERS


def _table(table: SlideTable) -> Table | None:
    rows = [["" if cell.is_spanned else _one_line(cell.text) for cell in row.cells] for row in table.rows]
    if not any(any(row) for row in rows):
        return None
    has_header = table.first_row and len(rows) >= _MIN_ROWS_FOR_A_HEADER  # PowerPoint's "Header Row" option
    return Table.from_ragged(rows, header_rows=1 if has_header else 0)


def _chart(chart: Chart) -> list[Block]:
    """A caption naming the chart, then the chart's own data as a table: categories down, one column per series."""
    title = _one_line(chart.chart_title.text_frame.text) if chart.has_title else ""
    caption = TextBlock(text=f"Chart: {title}" if title else "Chart", kind=TextKind.CAPTION)
    plots = list(chart.plots)
    categories = [str(category) for category in plots[0].categories] if plots else []
    series = list(plots[0].series) if plots else []
    if not categories or not series:  # XY and bubble charts have no categories
        return [caption]
    header = ["Category", *(item.name or f"Series {index}" for index, item in enumerate(series, start=1))]
    values = [tuple(item.values) for item in series]
    rows = [
        [category, *(_chart_value(column[row]) if row < len(column) else "" for column in values)]
        for row, category in enumerate(categories)
    ]
    return [caption, Table.from_ragged([header, *rows], header_rows=1)]


def _chart_value(value: float | None) -> str:
    """120.0 -> "120", 4.25 -> "4.25", missing -> ""."""
    if value is None:
        return ""
    rounded = float(f"{value:.{_CHART_DIGITS}g}")
    return str(int(rounded)) if rounded.is_integer() else str(rounded)


def _one_line(text: str) -> str:
    """Text with line breaks (PowerPoint's vertical tab) and runs of spaces collapsed to single spaces."""
    return " ".join(text.split())
