"""Document model -> Word with python-docx (IDP-03 ch. 3.4: ``model_to_docx``).

Headings, paragraphs, list items and captions use Word's own styles; tables are real grids whose header rows
repeat on every page. Heading styles are made plain black, as in the original ``to_docx.py`` (Sumukh, Karunya),
because Word's default theme makes them blue.
"""

import docx
from docx.document import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import RGBColor

from conversion.conversion.domain.document_model import Alignment, Table, TextBlock, TextKind, check_table_sizes
from conversion.conversion.domain.ports import StepRequest, StepResult
from conversion.conversion.domain.route import Converter
from conversion.conversion.infrastructure.model_file import read_model, source_stem

_BLACK = RGBColor(0, 0, 0)
_PLAIN_STYLES = ("Title", *(f"Heading {level}" for level in range(1, 7)))
_LIST_STYLES = ("List Bullet", "List Bullet 2", "List Bullet 3")  # the deepest Word's default template offers
_ALIGNMENTS = {
    Alignment.LEFT: WD_ALIGN_PARAGRAPH.LEFT,
    Alignment.CENTER: WD_ALIGN_PARAGRAPH.CENTER,
    Alignment.RIGHT: WD_ALIGN_PARAGRAPH.RIGHT,
}


class PythonDocxWriter:
    """Writes ``<original>.docx``. Each named page (a sheet or slide) starts with its name as a heading."""

    def __init__(self, *, max_table_rows: int) -> None:
        self._max_table_rows = max_table_rows

    @property
    def converter(self) -> Converter:
        return Converter.MODEL_TO_DOCX

    def convert(self, request: StepRequest) -> StepResult:
        (model_file,) = request.inputs
        model = read_model(model_file)
        check_table_sizes(model, self._max_table_rows)
        document = docx.Document()
        _make_headings_plain(document)
        for page in model.pages:
            if page.name:
                document.add_heading(page.name, level=1)
            for block in page.blocks:
                if isinstance(block, Table):
                    _add_table(document, block)
                else:
                    _add_text(document, block)
        output = request.output_dir / f"{source_stem(model_file)}.docx"
        document.save(str(output))
        return StepResult(outputs=(output,), page_count=max(len(model.pages), 1))


def _make_headings_plain(document: WordDocument) -> None:
    for name in _PLAIN_STYLES:
        document.styles[name].font.color.rgb = _BLACK


def _add_text(document: WordDocument, block: TextBlock) -> None:
    if block.kind is TextKind.HEADING:
        document.add_heading(block.text, level=block.level)
        return
    style = {
        TextKind.LIST_ITEM: _LIST_STYLES[min(block.level, len(_LIST_STYLES)) - 1],
        TextKind.CAPTION: "Caption",
    }.get(block.kind)
    paragraph = document.add_paragraph(block.text, style=style)
    paragraph.alignment = _ALIGNMENTS[block.alignment]
    if block.bold:
        for run in paragraph.runs:
            run.bold = True


def _add_table(document: WordDocument, table: Table) -> None:
    grid = document.add_table(rows=len(table.rows), cols=table.column_count)
    grid.style = "Table Grid"
    for index, (row, values) in enumerate(zip(grid.rows, table.rows, strict=True)):
        is_header = index < table.header_rows
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
            if is_header:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
        if is_header:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))  # repeat on each page
