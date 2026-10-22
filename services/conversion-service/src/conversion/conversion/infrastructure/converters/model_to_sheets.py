"""Document model -> XLSX or CSV (IDP-03 ch. 3.4: ``model_to_xlsx``, ``model_to_csv``).

Data only (fidelity D): tables are kept; headings and paragraphs are not. XLSX gets one sheet per table, typed
with the same ``cell_values`` rules as CSV -> XLSX. CSV gets one table: the first, or the first on the page the
``sheet`` option names.
"""

import csv

from openpyxl import Workbook

from conversion.conversion.domain.document_model import DocumentModel, Table
from conversion.conversion.domain.exceptions import NoTableFoundError, SheetNotFoundError
from conversion.conversion.domain.ports import StepRequest, StepResult
from conversion.conversion.domain.route import Converter
from conversion.conversion.infrastructure.converters.spreadsheets import append_table, sheet_title
from conversion.conversion.infrastructure.model_file import read_model, source_stem


class OpenpyxlModelWriter:
    """Writes ``<original>.xlsx`` with one sheet per table, named after its page when that page has one table."""

    @property
    def converter(self) -> Converter:
        return Converter.MODEL_TO_XLSX

    def convert(self, request: StepRequest) -> StepResult:
        (model_file,) = request.inputs
        model = read_model(model_file)
        tables = [
            (page.name if len(page.tables) == 1 else None, table) for page in model.pages for table in page.tables
        ]
        if not tables:
            raise NoTableFoundError
        workbook, taken = Workbook(write_only=True), set[str]()
        for number, (name, table) in enumerate(tables, start=1):
            sheet = workbook.create_sheet(sheet_title(name or f"Table {number}", taken))
            append_table(sheet, table.rows, table.header_rows)
        output = request.output_dir / f"{source_stem(model_file)}.xlsx"
        workbook.save(output)
        return StepResult(outputs=(output,), page_count=len(tables))


class CsvModelWriter:
    """Writes ``<original>.csv`` from one table, as UTF-8."""

    @property
    def converter(self) -> Converter:
        return Converter.MODEL_TO_CSV

    def convert(self, request: StepRequest) -> StepResult:
        (model_file,) = request.inputs
        table = choose_table(read_model(model_file), request.options.sheet)
        output = request.output_dir / f"{source_stem(model_file)}.csv"
        with output.open("w", encoding="utf-8", newline="") as stream:
            csv.writer(stream).writerows(table.rows)
        return StepResult(outputs=(output,), page_count=1)


def choose_table(model: DocumentModel, sheet: str | None) -> Table:
    """The first table, or the first on the page named ``sheet``."""
    if sheet is None:
        tables = model.tables
    else:
        page = next((page for page in model.pages if page.name == sheet), None)
        if page is None:
            raise SheetNotFoundError(sheet)
        tables = page.tables
    if not tables:
        raise NoTableFoundError
    return tables[0]
