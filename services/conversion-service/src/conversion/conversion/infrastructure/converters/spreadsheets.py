"""Reading and writing spreadsheets, shared by the CSV, XLSX and document-model converters.

CSV reading (encoding and delimiter detection) is ported from ``csv_convert.py`` (Sumukh, Karunya). Typed table
writing applies the ``cell_values`` rules, so CSV -> XLSX and document -> XLSX type numbers the same way.
"""

import csv
import io
import re
import zipfile
from collections.abc import Iterator, Sequence
from contextlib import contextmanager
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook as LoadedWorkbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.worksheet._write_only import WriteOnlyWorksheet

from conversion.conversion.domain.cell_values import ColumnKind, ColumnPlan, number_format, plan_columns
from conversion.conversion.domain.exceptions import LimitExceededError, SheetNotFoundError, UnreadableDocumentError

EXCEL_MAX_ROWS, EXCEL_MAX_COLUMNS = 1_048_576, 16_384
EXCEL_MAX_SHEET_TITLE = 31
_DELIMITERS = ",;\t|"
_ENCODINGS = ("utf-8-sig", "cp1252")
_LAST_RESORT = "latin-1"  # decodes any bytes, so it never fails
_UTF16_MARKS = (b"\xff\xfe", b"\xfe\xff")
_BINARY_PROBE_BYTES = 4096
_SNIFF_CHARS = 65_536
_TITLE_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")  # characters Excel refuses in sheet names
_TEXT_FORMAT = "@"  # Excel's "Text" format: what is written is what is shown
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F3D6D")
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
_RIGHT = Alignment(horizontal="right")
_MIN_WIDTH, _MAX_WIDTH, _WIDTH_SAMPLE_ROWS = 6, 60, 2000


# --- CSV ------------------------------------------------------------------------------------------------------


def read_csv(path: Path) -> list[list[str]]:
    """Rows of trimmed cells, empty rows dropped, every row padded to the widest one."""
    raw = path.read_bytes()  # bounded by the plan's upload size, enforced by document-service
    utf16 = raw.startswith(_UTF16_MARKS)
    if b"\x00" in raw[:_BINARY_PROBE_BYTES] and not utf16:
        raise UnreadableDocumentError  # a binary file renamed to .csv
    text = _decode(raw, utf16=utf16)
    try:
        reader = csv.reader(io.StringIO(text), delimiter=_sniff_delimiter(text[:_SNIFF_CHARS]), skipinitialspace=True)
        rows = [[cell.strip() for cell in row] for row in reader]
    except csv.Error as exc:
        raise UnreadableDocumentError from exc
    rows = [row for row in rows if any(row)]
    width = max((len(row) for row in rows), default=0)
    if len(rows) > EXCEL_MAX_ROWS or width > EXCEL_MAX_COLUMNS:
        raise LimitExceededError("The table is larger than a spreadsheet can hold (1,048,576 rows by 16,384 columns).")
    return [row + [""] * (width - len(row)) for row in rows]


def _decode(raw: bytes, *, utf16: bool) -> str:
    for encoding in (("utf-16",) if utf16 else ()) + _ENCODINGS:
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode(_LAST_RESORT)


def _sniff_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=_DELIMITERS).delimiter
    except csv.Error:
        counts = {delimiter: sample.count(delimiter) for delimiter in _DELIMITERS}
        best = max(counts, key=lambda delimiter: counts[delimiter])
        return best if counts[best] else ","


# --- Reading XLSX ---------------------------------------------------------------------------------------------


@contextmanager
def open_workbook(path: Path) -> Iterator[LoadedWorkbook]:
    """Open a workbook read-only, with cached formula results rather than formulas; always closed."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError) as exc:
        raise UnreadableDocumentError from exc
    try:
        yield workbook
    finally:
        workbook.close()


def pick_sheet(workbook: LoadedWorkbook, name: str | None) -> ReadOnlyWorksheet:
    """The named sheet, or the first. A chart-only sheet has no cells, so it doesn't count as a sheet here."""
    sheet = workbook.worksheets[0] if name is None else workbook[name] if name in workbook.sheetnames else None
    if not isinstance(sheet, ReadOnlyWorksheet):
        raise SheetNotFoundError(name or "")
    return sheet


def worksheet_rows(sheet: ReadOnlyWorksheet) -> Iterator[list[str]]:
    """Every row as text cells, streamed."""
    # ReadOnlyWorksheet borrows Worksheet.iter_rows at runtime; the stub types `self` as Worksheet.
    for row in sheet.iter_rows(values_only=True):  # type: ignore[misc]
        yield [as_text(value) for value in row]


def as_text(value: object) -> str:
    """A cell value as text: numbers without exponents, dates in ISO form, booleans as TRUE/FALSE."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float | Decimal):
        return _number_text(value)
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat(sep=" ")
    if isinstance(value, date | time):
        return value.isoformat()
    return str(value)


def _number_text(value: float | Decimal) -> str:
    exact = Decimal(repr(value)) if isinstance(value, float) else value  # repr: the shortest exact digits
    return str(int(exact)) if exact == exact.to_integral_value() else format(exact, "f")


# --- Writing XLSX ---------------------------------------------------------------------------------------------


def append_table(sheet: WriteOnlyWorksheet, rows: Sequence[Sequence[str]], header_rows: int) -> None:
    """Write styled header rows, then typed body cells. Column widths are set first (write-only sheets need that)."""
    if not rows:
        return
    body = rows[header_rows:]
    plans = plan_columns(body, len(rows[0]))
    _set_column_widths(sheet, rows)
    for row in rows[:header_rows]:
        sheet.append([_header_cell(sheet, value) for value in row])
    for index, row in enumerate(body):
        sheet.append([_data_cell(sheet, value, plan, index) for value, plan in zip(row, plans, strict=True)])


def sheet_title(name: str, taken: set[str]) -> str:
    """A sheet name Excel accepts: forbidden characters removed, at most 31 characters, unique in the workbook."""
    base = _TITLE_FORBIDDEN.sub(" ", name).strip()[:EXCEL_MAX_SHEET_TITLE] or "Sheet"
    title, number = base, 2
    while title.casefold() in taken:
        suffix = f" ({number})"
        title, number = base[: EXCEL_MAX_SHEET_TITLE - len(suffix)] + suffix, number + 1
    taken.add(title.casefold())
    return title


def _header_cell(sheet: WriteOnlyWorksheet, value: str) -> Cell:
    cell = _text_cell(sheet, value)
    cell.font, cell.fill, cell.alignment = _HEADER_FONT, _HEADER_FILL, _HEADER_ALIGNMENT
    return cell


def _data_cell(sheet: WriteOnlyWorksheet, value: str, plan: ColumnPlan, index: int) -> Cell:
    number = plan.numbers[index] if plan.kind is ColumnKind.NUMBER else None
    if number is None:
        cell = _text_cell(sheet, value)
        if value and plan.kind is not ColumnKind.TEXT:
            cell.alignment = _RIGHT  # a number kept as text still lines up like one
        return cell
    # The file format stores numbers as binary doubles; parse_number keeps at most 15 significant digits,
    # which a double holds exactly, so converting here loses nothing.
    cell = WriteOnlyCell(sheet, value=float(number.value))
    cell.number_format = number_format(number, column_grouped=plan.grouped)
    return cell


def _text_cell(sheet: WriteOnlyWorksheet, value: str) -> Cell:
    """A cell that always holds text. openpyxl would store "=..." as a live formula (formula injection)."""
    cell = WriteOnlyCell(sheet, value=value or None)
    if value:
        cell.data_type = "s"
        cell.number_format = _TEXT_FORMAT
    return cell


def _set_column_widths(sheet: WriteOnlyWorksheet, rows: Sequence[Sequence[str]]) -> None:
    for column in range(len(rows[0])):
        longest = max(len(row[column]) for row in rows[:_WIDTH_SAMPLE_ROWS])
        width = min(max(longest + 2, _MIN_WIDTH), _MAX_WIDTH)
        dimensions = sheet.column_dimensions  # type: ignore[attr-defined]  # missing from types-openpyxl; exists at runtime
        dimensions[get_column_letter(column + 1)].width = width
