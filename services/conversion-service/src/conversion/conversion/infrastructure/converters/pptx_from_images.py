"""Images -> PowerPoint, one slide per image (IDP-03 ch. 3.4: ``pptx_from_images``).

Used for PNG/JPEG -> PPTX, and after ``pdf_to_images`` for PDF and Word -> PPTX. Each image is embedded as it
is, byte for byte, and fitted to the slide without cropping. The deck takes the shape of the first image, so a
portrait PDF gives portrait slides. A photo whose EXIF tag says "rotate" is turned upright first, because
PowerPoint ignores that tag.
"""

import tempfile
from pathlib import Path

import pptx
from PIL import ImageOps

from conversion.conversion.domain.ports import StepRequest, StepResult
from conversion.conversion.domain.route import Converter
from conversion.conversion.infrastructure.converters.pillow_images import open_image

SLIDE_WIDTH_EMU = 12_192_000  # 13.33 in, PowerPoint's widescreen width
_MIN_SIDE_EMU, _MAX_SIDE_EMU = 914_400, 51_206_400  # PowerPoint allows slides from 1 to 56 inches
_EXIF_ORIENTATION = 0x0112
_BLANK_LAYOUT = 6  # in python-pptx's default template


class PptxImageSlides:
    """Writes one PPTX with one slide per input image, in input order, named after the first image."""

    def __init__(self, *, max_pixels: int) -> None:
        self._max_pixels = max_pixels

    @property
    def converter(self) -> Converter:
        return Converter.PPTX_FROM_IMAGES

    def convert(self, request: StepRequest) -> StepResult:
        deck = pptx.Presentation()
        layout = deck.slide_layouts[_BLANK_LAYOUT]
        with tempfile.TemporaryDirectory(prefix="upright-") as scratch:  # always cleaned up (FIL-06)
            for position, source in enumerate(request.inputs, start=1):
                picture, (width, height) = self._prepare(source, Path(scratch) / f"slide-{position:04d}.png")
                if position == 1:
                    deck.slide_width = pptx.util.Emu(SLIDE_WIDTH_EMU)
                    deck.slide_height = pptx.util.Emu(_slide_height(width, height))
                slide_width, slide_height = int(deck.slide_width or 0), int(deck.slide_height or 0)
                scale = min(slide_width / width, slide_height / height)
                shown_width, shown_height = round(width * scale), round(height * scale)
                deck.slides.add_slide(layout).shapes.add_picture(
                    str(picture),
                    pptx.util.Emu((slide_width - shown_width) // 2),
                    pptx.util.Emu((slide_height - shown_height) // 2),
                    pptx.util.Emu(shown_width),
                    pptx.util.Emu(shown_height),
                )
            output = request.output_dir / f"{request.inputs[0].stem}.pptx"
            deck.save(str(output))
        return StepResult(outputs=(output,), page_count=len(request.inputs))

    def _prepare(self, source: Path, upright: Path) -> tuple[Path, tuple[int, int]]:
        """Check ``source`` opens and isn't too big; return it, or an upright copy, with its size in pixels."""
        with open_image(source, max_pixels=self._max_pixels) as image:
            if int(image.getexif().get(_EXIF_ORIENTATION, 1) or 1) == 1:
                return source, image.size  # embedded unchanged
            turned = ImageOps.exif_transpose(image)
            turned.save(upright, format="PNG")
            return upright, turned.size


def _slide_height(width: int, height: int) -> int:
    """The slide height that gives the first image's shape at the standard width, within PowerPoint's limits."""
    return min(max(round(SLIDE_WIDTH_EMU * height / width), _MIN_SIDE_EMU), _MAX_SIDE_EMU)
