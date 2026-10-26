"""Document model -> PowerPoint with python-pptx (IDP-03 ch. 3.4: ``model_to_pptx``, "table slides, <= 15 rows").

A slide is a fixed canvas, so nothing is squeezed to fit: each page starts a slide titled with its name, and
content that doesn't fit continues on slides titled "<name> (cont.)", as in the original ``to_pptx.py``
(Karunya, Sumukh). A table breaks between rows with its header repeated; text breaks between blocks.
"""

from collections.abc import Iterator

import pptx
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.slide import Slide
from pptx.text.text import _Paragraph
from pptx.util import Emu, Length, Pt

from conversion.conversion.domain.document_model import (
    Alignment,
    Block,
    Page,
    Table,
    TextBlock,
    TextKind,
    check_table_sizes,
)
from conversion.conversion.domain.ports import StepRequest, StepResult
from conversion.conversion.domain.route import Converter
from conversion.conversion.infrastructure.model_file import read_model, source_stem

ROWS_PER_SLIDE = 15  # IDP-03 ch. 3.4; header rows included
TEXT_BLOCKS_PER_SLIDE = 10
_SLIDE_WIDTH, _SLIDE_HEIGHT = Emu(12_192_000), Emu(6_858_000)  # 13.33 x 7.5 in, 16:9; the default is 4:3
_MARGIN = Emu(457_200)  # 0.5 in
_TITLE_HEIGHT = Emu(1_005_840)  # 1.1 in
_BODY_TOP = Emu(_MARGIN + _TITLE_HEIGHT)
_BODY_WIDTH = Emu(_SLIDE_WIDTH - 2 * _MARGIN)
_BODY_HEIGHT = Emu(_SLIDE_HEIGHT - _BODY_TOP - _MARGIN)
_TABLE_ROW_HEIGHT = Emu(_BODY_HEIGHT // ROWS_PER_SLIDE)
_BULLET_INDENT = Emu(285_750)  # 0.3125 in per list level
_MAJOR_HEADING_LEVEL = 2  # headings at this level or above get the larger size
_TITLE_ONLY_LAYOUT = 5  # in python-pptx's default template
_TITLE_PLACEHOLDER = 0
_CELL_FONT_SIZES = ((5, Pt(14)), (8, Pt(12)), (12, Pt(10)))  # (up to this many columns, font size)
_SMALLEST_CELL_FONT = Pt(8)
_ALIGNMENTS = {Alignment.LEFT: PP_ALIGN.LEFT, Alignment.CENTER: PP_ALIGN.CENTER, Alignment.RIGHT: PP_ALIGN.RIGHT}

type _Body = tuple[TextBlock, ...] | Table


class PythonPptxWriter:
    """Writes ``<original>.pptx``: one or more slides per page."""

    def __init__(self, *, max_table_rows: int) -> None:
        self._max_table_rows = max_table_rows

    @property
    def converter(self) -> Converter:
        return Converter.MODEL_TO_PPTX

    def convert(self, request: StepRequest) -> StepResult:
        (model_file,) = request.inputs
        model = read_model(model_file)
        check_table_sizes(model, self._max_table_rows)
        stem = source_stem(model_file)
        deck = pptx.Presentation()
        deck.slide_width, deck.slide_height = _SLIDE_WIDTH, _SLIDE_HEIGHT
        layout = deck.slide_layouts[_TITLE_ONLY_LAYOUT]
        for page in model.pages or (Page(blocks=()),):
            title, blocks = _title(page, fallback=stem)
            for index, body in enumerate(list(_bodies(blocks)) or [()]):
                slide = deck.slides.add_slide(layout)
                _set_title(slide, title if index == 0 else f"{title} (cont.)")
                if isinstance(body, Table):
                    _add_table(slide, body)
                elif body:
                    _add_text(slide, body)
        output = request.output_dir / f"{stem}.pptx"
        deck.save(str(output))
        return StepResult(outputs=(output,), page_count=len(deck.slides))


def _title(page: Page, *, fallback: str) -> tuple[str, tuple[Block, ...]]:
    """The page name; else a leading heading, which then isn't repeated in the body; else the file name."""
    if page.name:
        return page.name, page.blocks
    first = page.blocks[0] if page.blocks else None
    if isinstance(first, TextBlock) and first.kind is TextKind.HEADING:
        return first.text, page.blocks[1:]
    return fallback, page.blocks


def _bodies(blocks: tuple[Block, ...]) -> Iterator[_Body]:
    """What goes on each slide: a run of up to ``TEXT_BLOCKS_PER_SLIDE`` text blocks, or one part of a table."""
    text: list[TextBlock] = []
    for block in blocks:
        if isinstance(block, Table):
            if text:
                yield tuple(text)
                text = []
            yield from _table_parts(block)
        else:
            text.append(block)
            if len(text) == TEXT_BLOCKS_PER_SLIDE:
                yield tuple(text)
                text = []
    if text:
        yield tuple(text)


def _table_parts(table: Table) -> Iterator[Table]:
    """The table cut into parts of at most ``ROWS_PER_SLIDE`` rows, each starting with the header rows."""
    header_rows = table.header_rows if table.header_rows < ROWS_PER_SLIDE else 0  # a header can't fill a slide
    header, body = table.rows[:header_rows], table.rows[header_rows:]
    room = ROWS_PER_SLIDE - header_rows
    for start in range(0, max(len(body), 1), room):
        yield Table(rows=header + body[start : start + room], header_rows=header_rows)


def _set_title(slide: Slide, text: str) -> None:
    title = slide.placeholders[_TITLE_PLACEHOLDER]  # placed for 4:3 by the template; moved to fit 16:9
    title.left, title.top, title.width, title.height = _MARGIN, _MARGIN // 2, _BODY_WIDTH, _TITLE_HEIGHT
    title.text_frame.text = text


def _add_text(slide: Slide, blocks: tuple[TextBlock, ...]) -> None:
    frame = slide.shapes.add_textbox(_MARGIN, _BODY_TOP, _BODY_WIDTH, _BODY_HEIGHT).text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE  # PowerPoint shrinks the text if it would overflow
    for index, block in enumerate(blocks):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = block.text
        paragraph.alignment = _ALIGNMENTS[block.alignment]
        heading = block.kind is TextKind.HEADING
        for run in paragraph.runs:
            run.font.size = _text_size(block)
            run.font.bold = block.bold or heading
            run.font.italic = block.kind is TextKind.CAPTION
        if block.kind is TextKind.LIST_ITEM:
            _make_bullet(paragraph, block.level)


def _text_size(block: TextBlock) -> Length:
    if block.kind is TextKind.HEADING:
        return Pt(22) if block.level <= _MAJOR_HEADING_LEVEL else Pt(18)
    return Pt(12) if block.kind is TextKind.CAPTION else Pt(16)


def _make_bullet(paragraph: _Paragraph, level: int) -> None:
    """A real bullet, indented by depth; python-pptx has no API for bullets in a text box."""
    paragraph.level = level - 1
    properties = paragraph._p.get_or_add_pPr()
    properties.set("marL", str(_BULLET_INDENT * level))
    properties.set("indent", str(-_BULLET_INDENT))
    properties.append(properties.makeelement(qn("a:buChar"), {"char": "•"}))


def _add_table(slide: Slide, table: Table) -> None:
    height = Emu(_TABLE_ROW_HEIGHT * len(table.rows))
    grid = slide.shapes.add_table(len(table.rows), table.column_count, _MARGIN, _BODY_TOP, _BODY_WIDTH, height).table
    grid.first_row = table.header_rows > 0  # the table style shows the first row as a header
    size = next((size for columns, size in _CELL_FONT_SIZES if table.column_count <= columns), _SMALLEST_CELL_FONT)
    for index, (row, values) in enumerate(zip(grid.rows, table.rows, strict=True)):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = size
                    run.font.bold = index < table.header_rows or None  # None: the table style decides
